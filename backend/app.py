import base64
import concurrent.futures
import hashlib
import html
import io
import json
import mimetypes
import os
import re
import secrets
import subprocess
import tempfile
import time
import urllib.parse
import zipfile
from collections import defaultdict, deque
from collections.abc import Generator
from datetime import datetime

import mistune
import psycopg
import requests
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.shared import RGBColor as DocxRGBColor
from dotenv import load_dotenv
from fastapi import (
    BackgroundTasks,
    FastAPI,
    File,
    Form,
    HTTPException,
    Request,
    UploadFile,
)
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import RedirectResponse, Response, StreamingResponse
from openai import OpenAI
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.util import Inches as PptxInches
from psycopg.errors import UniqueViolation
from psycopg.types.json import Json
from pydantic import BaseModel
from pypdf import PdfReader
from reportlab.lib import colors as reportlab_colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
import sentry_sdk
from sentry_sdk.integrations.fastapi import FastApiIntegration
from sentry_sdk.integrations.starlette import StarletteIntegration
from starlette.middleware.sessions import SessionMiddleware
from werkzeug.security import (
    check_password_hash,
    generate_password_hash,
)
from xhtml2pdf import pisa


load_dotenv()

# --- Error tracking (Phase 26) -----------------------------------------
# Optional, same "unset = skip gracefully" pattern as Google Sign-In
# below -- an unset SENTRY_DSN just means sentry_sdk.init() never runs.
# sentry_sdk itself is always imported (a hard dependency now -- see
# requirements.txt) so report_error() further down can call
# sentry_sdk.capture_exception() unconditionally: that's a documented
# safe no-op when init() was never called, so nothing needs to branch
# on SENTRY_DSN at every one of report_error()'s call sites, only here.
# This has to run before FastAPI() is instantiated below so the
# Starlette/FastAPI integration can actually instrument the app (it
# hooks in at app-creation time, not afterward).
SENTRY_DSN = os.getenv("SENTRY_DSN", "").strip()
SENTRY_ENVIRONMENT = os.getenv("SENTRY_ENVIRONMENT", "development")

if SENTRY_DSN:
    sentry_sdk.init(
        dsn=SENTRY_DSN,
        environment=SENTRY_ENVIRONMENT,
        # Conservative default (0 = no performance traces sent, only
        # errors) -- tracing has its own Sentry quota/cost separate
        # from error events, and this app has no need for it yet.
        traces_sample_rate=float(os.getenv("SENTRY_TRACES_SAMPLE_RATE", "0")),
        integrations=[StarletteIntegration(), FastApiIntegration()],
    )

app = FastAPI(title="EASTA API")


DATABASE_URL = os.getenv("DATABASE_URL")
OPENROUTER_API_KEY = os.getenv("OPENROUTER_API_KEY")

FRONTEND_URL = os.getenv(
    "FRONTEND_URL",
    "http://127.0.0.1:5000",
)

SESSION_SECRET = os.getenv(
    "SESSION_SECRET",
    "local-development-secret-change-me",
)

COOKIE_SECURE = (
    os.getenv("COOKIE_SECURE", "false").lower() == "true"
)

COOKIE_SAMESITE = os.getenv(
    "COOKIE_SAMESITE",
    "lax",
)

# --- Google Sign-In (OAuth 2.0 "Authorization Code" flow) ------------------
# Optional: unlike DATABASE_URL/OPENROUTER_API_KEY above, an unset
# GOOGLE_CLIENT_ID doesn't stop the app from starting -- GET
# /api/auth/google/login just 503s, and GET /api/features reports
# google_signin: false so the frontend hides the "Continue with Google"
# button instead of showing one that would just error (same pattern as
# the mic/TTS/research-mode feature flags below).
GOOGLE_CLIENT_ID = os.getenv("GOOGLE_CLIENT_ID", "")
GOOGLE_CLIENT_SECRET = os.getenv("GOOGLE_CLIENT_SECRET", "")
GOOGLE_REDIRECT_URI = os.getenv(
    "GOOGLE_REDIRECT_URI",
    "http://localhost:8000/api/auth/google/callback",
)
ENABLE_GOOGLE_SIGNIN = bool(GOOGLE_CLIENT_ID and GOOGLE_CLIENT_SECRET)

# --- Model routing config -------------------------------------------------
# EASTA picks between a fast/cheap model and a stronger "smart" model
# depending on how complex the incoming message looks, and falls back to
# a third model if the first choice errors out or is rate-limited.
FAST_MODEL = os.getenv(
    "EASTA_FAST_MODEL",
    "nvidia/nemotron-3-super-120b-a12b:free",
)
SMART_MODEL = os.getenv(
    "EASTA_SMART_MODEL",
    "openai/gpt-4.1-mini",
)
FALLBACK_MODEL = os.getenv(
    "EASTA_FALLBACK_MODEL",
    "meta-llama/llama-3.1-70b-instruct",
)

# Vision (image input) routing. EASTA_FAST_MODEL/EASTA_FALLBACK_MODEL
# above are not guaranteed to accept image input, so turns with an image
# attachment use this separate chain instead — see model_chain_for_images().
# Confirmed against OpenRouter's live catalog (https://openrouter.ai/models):
# "openai/gpt-4.1-mini" (the EASTA_SMART_MODEL default) already supports
# image input, so EASTA_VISION_MODEL defaults to it with no extra config;
# "google/gemini-3.8-flash" is a reasonably-priced, cross-provider vision
# fallback. Re-verify both against the catalog if you change either.
VISION_MODEL = os.getenv("EASTA_VISION_MODEL", SMART_MODEL)
VISION_FALLBACK_MODEL = os.getenv(
    "EASTA_VISION_FALLBACK_MODEL",
    "google/gemini-3.8-flash",
)

# Image generation, via OpenRouter's dedicated Image API (POST
# /api/v1/images — NOT the chat completions endpoint). Confirmed against
# OpenRouter's live catalog and docs; re-verify at
# https://openrouter.ai/models if you change it.
IMAGE_MODEL = os.getenv("EASTA_IMAGE_MODEL", "google/gemini-2.5-flash-image")

# Voice: server-side fallbacks for when the browser's native
# SpeechRecognition / speechSynthesis APIs aren't available (see
# ENABLE_SERVER_STT / ENABLE_SERVER_TTS below). Both go through
# OpenRouter's dedicated audio endpoints (POST /api/v1/audio/
# transcriptions and /api/v1/audio/speech) -- confirmed live and
# distinct from the chat completions endpoint used everywhere else.
STT_MODEL = os.getenv("EASTA_STT_MODEL", "openai/whisper-1")
TTS_MODEL = os.getenv("EASTA_TTS_MODEL", "openai/gpt-4o-mini-tts-2025-12-15")
TTS_VOICE = os.getenv("EASTA_TTS_VOICE", "alloy")

RATE_LIMIT_PER_MINUTE = int(
    os.getenv("EASTA_RATE_LIMIT_PER_MINUTE", "20")
)

# --- Feature flags ---------------------------------------------------------
ENABLE_TOOLS = (
    os.getenv("EASTA_ENABLE_TOOLS", "true").lower() == "true"
)
ENABLE_RAG = (
    os.getenv("EASTA_ENABLE_RAG", "true").lower() == "true"
)
ENABLE_SUMMARIZATION = (
    os.getenv("EASTA_ENABLE_SUMMARIZATION", "true").lower() == "true"
)
CONTEXT_RECENT_MESSAGES = int(
    os.getenv("EASTA_CONTEXT_RECENT_MESSAGES", "12")
)
# Cross-conversation memory (Phase 21): short, durable facts about a
# user extracted from what they explicitly say, injected into every
# future conversation's context -- see maybe_extract_memory() /
# build_memory_context_message() below.
ENABLE_MEMORY = (
    os.getenv("EASTA_ENABLE_MEMORY", "true").lower() == "true"
)
ENABLE_RESEARCH_MODE = (
    os.getenv("EASTA_ENABLE_RESEARCH_MODE", "true").lower() == "true"
)
ENABLE_GENERATION = (
    os.getenv("EASTA_ENABLE_GENERATION", "true").lower() == "true"
)
# Server-side voice fallbacks (see STT_MODEL/TTS_MODEL above). The
# composer mic button and each assistant message's "Read aloud" button
# use the browser's native APIs first and only need these as a
# fallback -- see GET /api/features, which the frontend uses to hide
# those controls entirely rather than show a control that would error.
ENABLE_SERVER_STT = (
    os.getenv("EASTA_ENABLE_SERVER_STT", "true").lower() == "true"
)
ENABLE_SERVER_TTS = (
    os.getenv("EASTA_ENABLE_SERVER_TTS", "true").lower() == "true"
)


def report_error(context: str, error: Exception) -> None:
    """Best-effort visibility for a caught, non-fatal failure that
    already doesn't interrupt the user's turn (background
    summarization, memory extraction, usage logging, ...) -- these used
    to only ever show up in stdout, easy to miss entirely outside of
    actively tailing server logs. Still prints (unchanged behavior for
    local dev with no Sentry configured), and additionally reports to
    Sentry when SENTRY_DSN is set. capture_exception() is a documented
    safe no-op when sentry_sdk.init() was never called, so this never
    needs to branch on SENTRY_DSN itself."""
    print(f"EASTA: {context}: {error!r}")
    sentry_sdk.capture_exception(error)


# Tool-calling is only attempted on models that are reasonably likely to
# support it. The free/fast model is skipped by default to keep simple
# turns cheap and low-latency — adjust this set as you validate specific
# OpenRouter models.
TOOL_CAPABLE_MODELS = {SMART_MODEL, FALLBACK_MODEL}

EASTA_SYSTEM_PROMPT = """\
You are EASTA, a helpful, knowledgeable AI assistant built to serve \
users across Africa and beyond. You are direct, accurate, and clear. \
You are fluent in English, French, Spanish, German, Chinese, \
Russian, Portuguese, Japanese, Korean, Italian, and Dutch. Always \
reply in the same language the user's latest message is written in, \
matching the register and formality norms native speakers expect in \
that language (for example vous/tu in French, formal/plain speech \
levels in Japanese and Korean, Sie/du in German) — unless the user \
asks you to reply in a different language, or a "Preferred reply \
language" system instruction is present, in which case follow that \
instruction instead. When you are unsure of a fact, say so instead \
of guessing. Keep responses well-organized and easy to read, using \
Markdown (headings, lists, code blocks) where it helps. You are not \
any underlying model provider's product — you are EASTA. \
When a "Grounding context" system message is present, prefer it over \
your own general knowledge and say so if it doesn't answer the \
question. When tools are available, use them only when they would \
materially improve the answer (e.g. current events, exact \
calculations) — do not narrate that you are "about to use a tool", \
just use it.\
"""

# Reply-language codes offered in the chat UI's language selector. "auto"
# means: no override, reply in whichever language the user's message is
# written in (the default instruction already baked into the system
# prompt above).
LANGUAGE_OPTIONS = {
    "auto": "Auto",
    "en": "English",
    "fr": "French",
    "es": "Spanish",
    "de": "German",
    "zh": "Chinese",
    "ru": "Russian",
    "pt": "Portuguese",
    "ja": "Japanese",
    "ko": "Korean",
    "it": "Italian",
    "nl": "Dutch",
}

if not DATABASE_URL:
    raise RuntimeError("DATABASE_URL is not configured.")

if not OPENROUTER_API_KEY:
    raise RuntimeError("OPENROUTER_API_KEY is not configured.")


client = OpenAI(
    base_url="https://openrouter.ai/api/v1",
    api_key=OPENROUTER_API_KEY,
)


app.add_middleware(
    SessionMiddleware,
    secret_key=SESSION_SECRET,
    session_cookie="easta_session",
    same_site=COOKIE_SAMESITE,
    https_only=COOKIE_SECURE,
)


app.add_middleware(
    CORSMiddleware,
    allow_origins=[FRONTEND_URL],
    allow_credentials=True,
    allow_methods=["GET", "POST", "PUT", "DELETE"],
    allow_headers=["Content-Type"],
)


class LoginRequest(BaseModel):
    username: str
    password: str


class RegisterRequest(BaseModel):
    username: str
    email: str
    password: str
    confirm_password: str


class AccountUpdateRequest(BaseModel):
    username: str
    email: str


class PasswordChangeRequest(BaseModel):
    current_password: str
    new_password: str
    confirm_new_password: str


class ApiKeyCreateRequest(BaseModel):
    name: str


class ConversationUpdateRequest(BaseModel):
    """All fields optional -- PATCH /api/conversations/{id} only
    touches whichever ones the client actually sent (see
    update_conversation_fields() and its exclude_unset=True caller),
    so a rename request doesn't accidentally clear pinned/folder and
    vice versa."""
    title: str | None = None
    pinned: bool | None = None
    folder: str | None = None


class PublicChatRequest(BaseModel):
    message: str
    conversation_id: int | None = None
    language: str | None = None


class ImageAttachment(BaseModel):
    name: str | None = None
    mime_type: str
    # A full data: URL, e.g. "data:image/png;base64,...." — see
    # validate_images() for size/type limits.
    data_url: str


class MessageRequest(BaseModel):
    message: str
    # Reply-language override from the chat UI's language selector (a
    # key from LANGUAGE_OPTIONS, e.g. "fr"), or "auto"/omitted for no
    # override.
    language: str | None = None
    images: list[ImageAttachment] = []
    # Research mode toggle from the composer — see run_research().
    research: bool = False
    # Composer image-style control (a key from IMAGE_ASPECT_RATIOS,
    # e.g. "portrait"), or "auto"/omitted to let the model choose via
    # generate_image's own aspect_ratio argument.
    image_aspect_ratio: str | None = None


class RegenerateRequest(BaseModel):
    language: str | None = None
    research: bool = False
    image_aspect_ratio: str | None = None


class DocumentRequest(BaseModel):
    title: str
    content: str


class CreateOrganizationRequest(BaseModel):
    name: str


class SpeakRequest(BaseModel):
    text: str
    # Reply-language code (see LANGUAGE_OPTIONS) -- currently unused by
    # synthesize_speech_bytes() since TTS_VOICE is a single fixed voice,
    # but accepted so the client can send it once per-language voice
    # selection is worth adding (see README "Suggested next steps").
    language: str | None = None


class RegenerateImageRequest(BaseModel):
    prompt: str
    aspect_ratio: str | None = None
    conversation_id: int | None = None


def require_user(request: Request) -> int:
    """Resolves the authenticated user id from either an
    `Authorization: Bearer <api key>` header (see resolve_api_key()) or
    the session cookie -- every existing require_user()-gated endpoint
    picks up Bearer-token auth for free from this one change, no
    per-route duplication needed. Both paths return the same kind of
    user_id, so enforce_rate_limit(user_id) downstream already applies
    equally to API-key traffic as to session traffic."""
    authorization = request.headers.get("Authorization", "")

    if authorization.startswith("Bearer "):
        api_key = authorization[len("Bearer "):].strip()
        user_id = resolve_api_key(api_key) if api_key else None

        if user_id is None:
            raise HTTPException(
                status_code=401,
                detail="Invalid or revoked API key.",
            )

        return user_id

    user_id = request.session.get("user_id")

    if user_id is None:
        raise HTTPException(
            status_code=401,
            detail="Please log in again.",
        )

    return user_id


# --- Simple in-memory rate limiter -----------------------------------------
# NOTE: this is per-process, so it resets on deploy and does not share state
# across multiple horizontally-scaled instances. Good enough for a single
# backend instance; swap for a Redis-backed limiter (e.g. via `slowapi` +
# Redis) once you scale the backend beyond one instance.
_request_log: dict[int, deque] = defaultdict(deque)


def enforce_rate_limit(user_id: int) -> None:
    now = time.monotonic()
    window = _request_log[user_id]

    while window and now - window[0] > 60:
        window.popleft()

    if len(window) >= RATE_LIMIT_PER_MINUTE:
        raise HTTPException(
            status_code=429,
            detail=(
                "You're sending messages too quickly. "
                "Please wait a moment and try again."
            ),
        )

    window.append(now)


# --- Auth / user helpers ----------------------------------------------------

def get_user_by_username(username: str):
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT
                    id,
                    username,
                    email,
                    password_hash
                FROM users
                WHERE username = %s;
                """,
                (username,),
            )

            return cursor.fetchone()


def get_user_by_email(email: str):
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT id
                FROM users
                WHERE email = %s;
                """,
                (email,),
            )

            return cursor.fetchone()


def create_user(
    username: str,
    email: str,
    password: str,
):
    password_hash = generate_password_hash(password)

    try:
        with psycopg.connect(DATABASE_URL) as connection:
            with connection.cursor() as cursor:
                cursor.execute(
                    """
                    INSERT INTO users (
                        username,
                        email,
                        password_hash
                    )
                    VALUES (%s, %s, %s)
                    RETURNING
                        id,
                        username,
                        email;
                    """,
                    (
                        username,
                        email,
                        password_hash,
                    ),
                )

                return cursor.fetchone()

    except UniqueViolation as error:
        raise HTTPException(
            status_code=409,
            detail=(
                "That username or email is already registered."
            ),
        ) from error


def get_user_by_google_id(google_id: str):
    """Returns (id, username, email) for session-cookie purposes --
    used on every Google sign-in after the first to find the existing
    account without re-matching by email."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT id, username, email
                FROM users
                WHERE google_id = %s;
                """,
                (google_id,),
            )

            return cursor.fetchone()


def link_google_id(user_id: int, google_id: str):
    """Attaches a Google account to an existing password-based user,
    matched by verified email in the callback below -- the account
    keeps its password (auth_provider stays 'password') and gains the
    ability to also log in with Google going forward."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE users
                SET google_id = %s
                WHERE id = %s
                RETURNING id, username, email;
                """,
                (google_id, user_id),
            )

            return cursor.fetchone()


_USERNAME_SANITIZE_RE = re.compile(r"[^a-zA-Z0-9_]")


def generate_unique_username(seed: str) -> str:
    """Derives a users.username (50 chars max, unique) from a Google
    display name or email local-part -- appends a numeric suffix on
    collision rather than failing the sign-in over a taken username the
    person never chose themselves."""
    cleaned = _USERNAME_SANITIZE_RE.sub("", seed).strip("_") or "user"
    candidate = cleaned[:50]
    suffix = 0

    while get_user_by_username(candidate):
        suffix += 1
        suffix_text = f"_{suffix}"
        candidate = f"{cleaned[: 50 - len(suffix_text)]}{suffix_text}"

    return candidate


def create_google_user(google_id: str, email: str, username: str):
    """Creates a Google-only account: no password_hash,
    auth_provider = 'google'. Returns (id, username, email)."""
    try:
        with psycopg.connect(DATABASE_URL) as connection:
            with connection.cursor() as cursor:
                cursor.execute(
                    """
                    INSERT INTO users (
                        username,
                        email,
                        password_hash,
                        google_id,
                        auth_provider
                    )
                    VALUES (%s, %s, NULL, %s, 'google')
                    RETURNING
                        id,
                        username,
                        email;
                    """,
                    (username, email, google_id),
                )

                return cursor.fetchone()

    except UniqueViolation as error:
        raise HTTPException(
            status_code=409,
            detail="An account with that email already exists.",
        ) from error


def get_user_by_id(user_id: int):
    """Returns (id, username, email, password_hash, created_at, plan,
    is_admin) for the Account page -- unlike get_user_by_username()/
    get_user_by_email() (used only for login/uniqueness checks), this
    includes everything the profile/plan/admin sections need in one
    query. is_admin appended at the end (not inserted in the middle)
    so every existing positional unpack of this tuple keeps working
    unless explicitly updated to read it."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT id, username, email, password_hash, created_at, plan, is_admin
                FROM users
                WHERE id = %s;
                """,
                (user_id,),
            )

            return cursor.fetchone()


def is_user_admin(user_id: int) -> bool:
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT is_admin FROM users WHERE id = %s;",
                (user_id,),
            )

            row = cursor.fetchone()
            return bool(row and row[0])


def require_admin(user_id: int) -> None:
    """403, not 404 -- unlike require_org_member()'s org-hiding 404,
    there's nothing to hide the existence of here (every user already
    knows /api/admin/* exists), so the ordinary "you can't do this"
    status is the right one."""
    if not is_user_admin(user_id):
        raise HTTPException(
            status_code=403,
            detail="Admin access required.",
        )


def update_user_profile(user_id: int, username: str, email: str):
    """Returns (id, username, email, created_at, plan)."""
    try:
        with psycopg.connect(DATABASE_URL) as connection:
            with connection.cursor() as cursor:
                cursor.execute(
                    """
                    UPDATE users
                    SET username = %s, email = %s
                    WHERE id = %s
                    RETURNING id, username, email, created_at, plan;
                    """,
                    (username, email, user_id),
                )

                return cursor.fetchone()

    except UniqueViolation as error:
        raise HTTPException(
            status_code=409,
            detail="That username or email is already taken.",
        ) from error


def update_user_password(user_id: int, password_hash: str):
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE users
                SET password_hash = %s
                WHERE id = %s;
                """,
                (password_hash, user_id),
            )


# --- API keys ------------------------------------------------------------
# Backs the Account page's "API Keys" section and Bearer-token auth for
# the public /v1/* API (see require_user() below). The plaintext key is
# generated here and returned to the caller exactly once; only its hash
# is ever persisted.
#
# NOTE on hashing: unlike users.password_hash (werkzeug's
# generate_password_hash, deliberately slow to resist brute-forcing a
# low-entropy human password), API keys are checked on *every*
# authenticated request rather than once at login, and are already
# high-entropy random tokens (256 bits from secrets.token_urlsafe) --
# brute-forcing one is infeasible regardless of hash speed. A fast
# SHA-256 is the right tool here (the same approach GitHub/Stripe-style
# API keys use), not a deliberately slow password hash on every request.

def generate_api_key() -> str:
    return f"easta_{secrets.token_urlsafe(32)}"


def hash_api_key(api_key: str) -> str:
    return hashlib.sha256(api_key.encode("utf-8")).hexdigest()


def create_api_key(user_id: int, name: str, key_hash: str):
    """Returns (id, name, created_at, last_used_at, revoked_at)."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO api_keys (user_id, name, key_hash)
                VALUES (%s, %s, %s)
                RETURNING id, name, created_at, last_used_at, revoked_at;
                """,
                (user_id, name, key_hash),
            )

            return cursor.fetchone()


def get_api_keys(user_id: int):
    """Returns rows of (id, name, created_at, last_used_at,
    revoked_at), newest first. Never includes key_hash."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT id, name, created_at, last_used_at, revoked_at
                FROM api_keys
                WHERE user_id = %s
                ORDER BY created_at DESC;
                """,
                (user_id,),
            )

            return cursor.fetchall()


def get_active_api_key_by_hash(key_hash: str):
    """Returns (id, user_id) for a non-revoked key, or None."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT id, user_id
                FROM api_keys
                WHERE key_hash = %s AND revoked_at IS NULL;
                """,
                (key_hash,),
            )

            return cursor.fetchone()


def touch_api_key(key_id: int):
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE api_keys
                SET last_used_at = CURRENT_TIMESTAMP
                WHERE id = %s;
                """,
                (key_id,),
            )


def revoke_api_key(key_id: int, user_id: int):
    """Returns the key's id if a matching, not-already-revoked key was
    found and revoked, else None."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE api_keys
                SET revoked_at = CURRENT_TIMESTAMP
                WHERE id = %s AND user_id = %s AND revoked_at IS NULL
                RETURNING id;
                """,
                (key_id, user_id),
            )

            row = cursor.fetchone()
            return row[0] if row else None


def resolve_api_key(api_key: str) -> int | None:
    """Looks up a Bearer-token API key and, if valid and not revoked,
    records its use and returns the owning user_id. Returns None for
    any invalid/unknown/revoked key -- callers treat that as
    unauthenticated, the same as a missing session."""
    key_hash = hash_api_key(api_key)
    row = get_active_api_key_by_hash(key_hash)

    if not row:
        return None

    key_id, user_id = row

    try:
        touch_api_key(key_id)
    except Exception as error:  # noqa: BLE001 - best-effort, auth still succeeds
        report_error("api key last_used_at update failed", error)

    return user_id


# --- Cross-conversation memory (Phase 21) -----------------------------------
# Backs the /account "Remembered" list (view/delete individual/clear all)
# and build_memory_context_message() above -- see database/schema.sql's
# user_memory table for the extraction/retention reasoning.

def save_user_memory(
    user_id: int, content: str, source_conversation_id: int | None
):
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO user_memory (
                    user_id, content, source_conversation_id
                )
                VALUES (%s, %s, %s)
                RETURNING id;
                """,
                (user_id, content, source_conversation_id),
            )

            return cursor.fetchone()[0]


def get_user_memory(user_id: int, limit: int | None = None):
    """Returns (id, content, source_conversation_id, created_at) rows,
    most recent first."""
    query = """
        SELECT id, content, source_conversation_id, created_at
        FROM user_memory
        WHERE user_id = %s
        ORDER BY created_at DESC
    """
    params = [user_id]

    if limit:
        query += " LIMIT %s"
        params.append(limit)

    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(query + ";", params)

            return cursor.fetchall()


def delete_user_memory_item(memory_id: int, user_id: int) -> bool:
    """Ownership check is inline in the WHERE clause (unlike the
    conversations get-then-mutate pattern) since there's no separate
    read step needed before deleting one memory item. Returns whether
    a row was actually deleted."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                DELETE FROM user_memory
                WHERE id = %s AND user_id = %s
                RETURNING id;
                """,
                (memory_id, user_id),
            )

            return cursor.fetchone() is not None


def clear_user_memory(user_id: int):
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                "DELETE FROM user_memory WHERE user_id = %s;",
                (user_id,),
            )


# --- Organizations (Phase 22) -----------------------------------------------
# Shared team workspaces. See database/schema.sql's organizations /
# organization_members comment for the full design, including the
# explicit "conversations stay private, only the knowledge base and
# billing are shared" decision. Route-level guards below
# (require_org_member() / require_org_owner()) mirror require_user()'s
# style; membership checks return 404 (not 403) for a non-member so a
# stranger can't even confirm an org id exists.

def create_organization(name: str, owner_user_id: int):
    """Creates the org and adds its creator as 'owner' in the same
    transaction -- an org can never exist with zero members. Returns
    (id, name, invite_token, created_at)."""
    invite_token = secrets.token_urlsafe(32)

    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO organizations (name, invite_token)
                VALUES (%s, %s)
                RETURNING id, name, invite_token, created_at;
                """,
                (name, invite_token),
            )
            org = cursor.fetchone()

            cursor.execute(
                """
                INSERT INTO organization_members (org_id, user_id, role)
                VALUES (%s, %s, 'owner');
                """,
                (org[0], owner_user_id),
            )

    return org


def get_user_organizations(user_id: int):
    """Returns (org_id, name, role) for every org this user belongs
    to, oldest membership first."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT o.id, o.name, om.role
                FROM organization_members om
                JOIN organizations o ON o.id = om.org_id
                WHERE om.user_id = %s
                ORDER BY om.joined_at;
                """,
                (user_id,),
            )

            return cursor.fetchall()


def get_organization_membership(org_id: int, user_id: int):
    """Returns this user's role ('owner'/'member') in this org, or
    None if they're not a member."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT role
                FROM organization_members
                WHERE org_id = %s AND user_id = %s;
                """,
                (org_id, user_id),
            )

            row = cursor.fetchone()
            return row[0] if row else None


def require_org_member(org_id: int, user_id: int) -> str:
    role = get_organization_membership(org_id, user_id)

    if not role:
        raise HTTPException(
            status_code=404,
            detail="Organization not found.",
        )

    return role


def require_org_owner(org_id: int, user_id: int) -> None:
    role = require_org_member(org_id, user_id)

    if role != "owner":
        raise HTTPException(
            status_code=403,
            detail="Only an organization owner can do this.",
        )


def get_organization(org_id: int):
    """Returns (id, name, invite_token, created_at) or None."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT id, name, invite_token, created_at
                FROM organizations
                WHERE id = %s;
                """,
                (org_id,),
            )

            return cursor.fetchone()


def get_organization_by_invite_token(token: str):
    """Returns (id, name) or None."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT id, name
                FROM organizations
                WHERE invite_token = %s;
                """,
                (token,),
            )

            return cursor.fetchone()


def regenerate_organization_invite_token(org_id: int) -> str:
    """Rotates the invite link -- same "revoke by changing the value"
    idea as conversations.share_token: any link built from the old
    token immediately stops working."""
    new_token = secrets.token_urlsafe(32)

    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE organizations
                SET invite_token = %s
                WHERE id = %s;
                """,
                (new_token, org_id),
            )

    return new_token


def join_organization(org_id: int, user_id: int, role: str = "member"):
    """Idempotent: joining an org you already belong to is a no-op,
    not an error."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO organization_members (org_id, user_id, role)
                VALUES (%s, %s, %s)
                ON CONFLICT (org_id, user_id) DO NOTHING;
                """,
                (org_id, user_id, role),
            )


def get_organization_members(org_id: int):
    """Returns (user_id, username, email, role, joined_at) rows."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT u.id, u.username, u.email, om.role, om.joined_at
                FROM organization_members om
                JOIN users u ON u.id = om.user_id
                WHERE om.org_id = %s
                ORDER BY om.joined_at;
                """,
                (org_id,),
            )

            return cursor.fetchall()


def count_organization_owners(org_id: int) -> int:
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT COUNT(*)
                FROM organization_members
                WHERE org_id = %s AND role = 'owner';
                """,
                (org_id,),
            )

            return cursor.fetchone()[0]


def remove_organization_member(org_id: int, user_id: int):
    """Caller (the route) is responsible for the "don't strand the org
    without an owner" check -- see count_organization_owners()."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                DELETE FROM organization_members
                WHERE org_id = %s AND user_id = %s;
                """,
                (org_id, user_id),
            )


def get_organization_usage(org_id: int):
    """Returns (user_id, username, total_cost, call_count) per member
    -- backs the owner-only combined cost dashboard. LEFT JOIN so a
    member with zero calls still shows up with 0/0 rather than being
    silently absent from their own org's breakdown."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT
                    u.id,
                    u.username,
                    COALESCE(SUM(ul.cost_usd), 0) AS total_cost,
                    COUNT(ul.id) AS call_count
                FROM organization_members om
                JOIN users u ON u.id = om.user_id
                LEFT JOIN usage_logs ul ON ul.user_id = u.id
                WHERE om.org_id = %s
                GROUP BY u.id, u.username
                ORDER BY total_cost DESC;
                """,
                (org_id,),
            )

            return cursor.fetchall()


# --- Documents (RAG knowledge base) -----------------------------------------
# Phase 22: a document belongs to either a user (personal) or an org
# (shared with every member) -- see database/schema.sql's
# documents_owner_check. Exactly one of user_id/org_id is passed to
# each helper below; the caller (the route) decides which, after its
# own auth/membership check.

def create_document_row(
    title: str,
    content: str,
    user_id: int | None = None,
    org_id: int | None = None,
):
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO documents (
                    user_id, org_id, title, content, search_vector
                )
                VALUES (
                    %s, %s, %s, %s, to_tsvector('english', %s)
                )
                RETURNING id, title, created_at;
                """,
                (user_id, org_id, title, content, content),
            )

            return cursor.fetchone()


def get_document_rows(user_id: int | None = None, org_id: int | None = None):
    """Returns (id, title, created_at, length) rows for exactly one of
    user_id/org_id."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            if org_id is not None:
                cursor.execute(
                    """
                    SELECT id, title, created_at, length(content)
                    FROM documents
                    WHERE org_id = %s
                    ORDER BY created_at DESC;
                    """,
                    (org_id,),
                )
            else:
                cursor.execute(
                    """
                    SELECT id, title, created_at, length(content)
                    FROM documents
                    WHERE user_id = %s
                    ORDER BY created_at DESC;
                    """,
                    (user_id,),
                )

            return cursor.fetchall()


def delete_document_row(
    document_id: int,
    user_id: int | None = None,
    org_id: int | None = None,
) -> bool:
    """Exactly one of user_id/org_id scopes the delete. For an org
    document, any member may delete it (a shared, collaboratively
    editable knowledge base, not owner-gated) -- the route has already
    confirmed the caller is a member before calling this. Returns
    whether a row was actually deleted."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            if org_id is not None:
                cursor.execute(
                    """
                    DELETE FROM documents
                    WHERE id = %s AND org_id = %s
                    RETURNING id;
                    """,
                    (document_id, org_id),
                )
            else:
                cursor.execute(
                    """
                    DELETE FROM documents
                    WHERE id = %s AND user_id = %s
                    RETURNING id;
                    """,
                    (document_id, user_id),
                )

            return cursor.fetchone() is not None


# --- Conversation / message helpers -----------------------------------------

def create_conversation(user_id: int):
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO conversations (user_id)
                VALUES (%s)
                RETURNING
                    id,
                    title,
                    created_at;
                """,
                (user_id,),
            )

            return cursor.fetchone()


def get_conversation(
    conversation_id: int,
    user_id: int,
):
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT
                    id,
                    title,
                    created_at,
                    summary,
                    summarized_through_message_id,
                    pinned,
                    folder,
                    share_token,
                    shared_at
                FROM conversations
                WHERE
                    id = %s
                    AND user_id = %s;
                """,
                (
                    conversation_id,
                    user_id,
                ),
            )

            return cursor.fetchone()


def set_conversation_share_token(conversation_id: int, token: str):
    """Caller is responsible for the ownership check
    (get_conversation(conversation_id, user_id)) before calling this.
    Setting a token when one already exists (re-sharing) overwrites
    it, which is exactly what "regenerating invalidates old links"
    means -- the old token string just stops matching any row."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE conversations
                SET share_token = %s, shared_at = CURRENT_TIMESTAMP
                WHERE id = %s;
                """,
                (token, conversation_id),
            )


def clear_conversation_share_token(conversation_id: int):
    """Caller is responsible for the ownership check first. Clearing
    both columns (not just share_token) so a later re-share shows a
    fresh "not shared before" state rather than a stale shared_at."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE conversations
                SET share_token = NULL, shared_at = NULL
                WHERE id = %s;
                """,
                (conversation_id,),
            )


def get_conversation_by_share_token(token: str):
    """Backs the unauthenticated GET /api/shared/{token} -- the whole
    point of this query is that it can ONLY ever return what a
    stranger with just the link is allowed to see, so it deliberately
    selects title only, never user_id/email/username or anything else
    about the account that shared it."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT id, title
                FROM conversations
                WHERE share_token = %s;
                """,
                (token,),
            )

            return cursor.fetchone()


def get_conversations(user_id: int):
    """Returns (id, title, created_at, chat_number, pinned, folder,
    share_token). chat_number is a stable chronological identity
    (oldest = 1, used for the "Chat N - date" auto-label) computed
    independently of display order, so pinning/unpinning a
    conversation never changes its number -- only the final ORDER BY
    (pinned first) changes where it renders in the sidebar."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT
                    id,
                    title,
                    created_at,
                    chat_number,
                    pinned,
                    folder,
                    share_token
                FROM (
                    SELECT
                        id,
                        title,
                        created_at,
                        pinned,
                        folder,
                        share_token,
                        ROW_NUMBER() OVER (
                            ORDER BY created_at, id
                        ) AS chat_number
                    FROM conversations
                    WHERE user_id = %s
                ) AS numbered_conversations
                ORDER BY pinned DESC, created_at DESC, id DESC;
                """,
                (user_id,),
            )

            return cursor.fetchall()


def rename_conversation_if_default(
    conversation_id: int,
    title: str,
):
    """Auto-title a conversation from the first user message, once."""
    trimmed = title.strip()[:100]

    if not trimmed:
        return

    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE conversations
                SET title = %s
                WHERE id = %s AND title = 'New Chat';
                """,
                (trimmed, conversation_id),
            )


_CONVERSATION_UPDATABLE_COLUMNS = {"title", "pinned", "folder"}


def update_conversation_fields(conversation_id: int, fields: dict):
    """Partial update for conversations.title/pinned/folder -- used by
    PATCH /api/conversations/{id} for rename (explicit user-driven,
    unlike rename_conversation_if_default() above which only touches
    the still-default title) and for pin/folder assignment alike, so a
    request only needs to send the field(s) it's actually changing.
    Caller is responsible for the ownership check
    (get_conversation(conversation_id, user_id)) before calling this,
    same pattern as every other write helper in this section.

    `fields` keys are filtered against the fixed whitelist above before
    ever reaching the query -- they're not trusted as literal SQL
    identifiers just because they originated from a Pydantic model."""
    columns = [key for key in fields if key in _CONVERSATION_UPDATABLE_COLUMNS]

    if not columns:
        return

    set_clause = ", ".join(f"{column} = %s" for column in columns)
    values = [fields[column] for column in columns] + [conversation_id]

    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                f"UPDATE conversations SET {set_clause} WHERE id = %s;",
                values,
            )


def delete_conversation(conversation_id: int):
    """Caller is responsible for the ownership check
    (get_conversation(conversation_id, user_id)) before calling this.
    messages/generated_files/canvas_artifacts all have ON DELETE
    CASCADE on their conversation_id foreign key (see
    database/schema.sql), and usage_logs has ON DELETE SET NULL (kept
    for historical cost accounting rather than deleted) -- nothing
    else to clean up manually here."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                DELETE FROM conversations
                WHERE id = %s;
                """,
                (conversation_id,),
            )


# Delimiters ts_headline() wraps each match in -- deliberately control
# characters that would never appear in real message/title text, so
# search_conversations() below can safely html.escape() the WHOLE
# snippet first (protecting against a user's own message content
# containing literal '<'/'>'/'&') and only THEN swap these markers for
# real <mark> tags, rather than trusting ts_headline's raw output
# (which defaults to <b>/</b> with no escaping of the surrounding
# text) directly into the page.
_SEARCH_HIGHLIGHT_START = "\x01"
_SEARCH_HIGHLIGHT_STOP = "\x02"
_TS_HEADLINE_OPTIONS = (
    f"StartSel={_SEARCH_HIGHLIGHT_START}, StopSel={_SEARCH_HIGHLIGHT_STOP}, "
    "MaxFragments=1, MaxWords=20, MinWords=5, HighlightAll=false"
)


def search_conversations(query: str, user_id: int, limit: int = 20):
    """Full-text search across a user's own conversation titles and
    message content (Postgres to_tsvector/websearch_to_tsquery, same
    technique as documents/RAG's retrieve_relevant_chunks() -- not a
    LIKE scan, so this stays fast as history grows). One result per
    conversation (its single best-ranked match, title or message),
    ordered by relevance. Returns (id, title, created_at, snippet,
    match_type) rows -- snippet is pre-escaped HTML with the match
    wrapped in <mark>, safe to insert directly."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            try:
                cursor.execute(
                    f"""
                    WITH query AS (
                        SELECT websearch_to_tsquery('english', %s) AS tsq
                    ),
                    title_matches AS (
                        SELECT
                            c.id,
                            c.title,
                            c.created_at,
                            ts_headline(
                                'english', c.title, query.tsq,
                                '{_TS_HEADLINE_OPTIONS}'
                            ) AS snippet,
                            'title' AS match_type,
                            ts_rank_cd(
                                to_tsvector('english', c.title), query.tsq
                            ) AS rank
                        FROM conversations c, query
                        WHERE
                            c.user_id = %s
                            AND to_tsvector('english', c.title) @@ query.tsq
                    ),
                    message_matches AS (
                        SELECT
                            c.id,
                            c.title,
                            c.created_at,
                            ts_headline(
                                'english', m.content, query.tsq,
                                '{_TS_HEADLINE_OPTIONS}'
                            ) AS snippet,
                            'message' AS match_type,
                            ts_rank_cd(m.search_vector, query.tsq) AS rank
                        FROM messages m
                        JOIN conversations c ON c.id = m.conversation_id
                        CROSS JOIN query
                        WHERE
                            c.user_id = %s
                            AND m.search_vector @@ query.tsq
                    ),
                    combined AS (
                        SELECT * FROM title_matches
                        UNION ALL
                        SELECT * FROM message_matches
                    ),
                    deduped AS (
                        SELECT DISTINCT ON (id) *
                        FROM combined
                        ORDER BY id, rank DESC
                    )
                    SELECT id, title, created_at, snippet, match_type
                    FROM deduped
                    ORDER BY rank DESC
                    LIMIT %s;
                    """,
                    (query, user_id, user_id, limit),
                )

                rows = cursor.fetchall()

            except psycopg.Error:
                # websearch_to_tsquery can reject odd input (e.g. bare
                # punctuation); degrade to "no results" instead of a
                # 500 for a query the user just typed oddly.
                connection.rollback()
                return []

    results = []
    for conversation_id, title, created_at, raw_snippet, match_type in rows:
        safe_snippet = (
            html.escape(raw_snippet)
            .replace(_SEARCH_HIGHLIGHT_START, "<mark>")
            .replace(_SEARCH_HIGHLIGHT_STOP, "</mark>")
        )
        results.append(
            (conversation_id, title, created_at, safe_snippet, match_type)
        )

    return results


def update_conversation_summary(
    conversation_id: int,
    summary: str,
    summarized_through_message_id: int,
):
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE conversations
                SET
                    summary = %s,
                    summarized_through_message_id = %s
                WHERE id = %s;
                """,
                (
                    summary,
                    summarized_through_message_id,
                    conversation_id,
                ),
            )


def save_message(
    conversation_id: int,
    role: str,
    content: str,
    attachments: list[dict] | None = None,
):
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO messages (
                    conversation_id,
                    role,
                    content,
                    attachments,
                    search_vector
                )
                VALUES (%s, %s, %s, %s, to_tsvector('english', %s))
                RETURNING id;
                """,
                (
                    conversation_id,
                    role,
                    content,
                    Json(attachments) if attachments else None,
                    content,
                ),
            )

            return cursor.fetchone()[0]


def get_messages(conversation_id: int):
    """Returns rows of (role, content, created_at, id, attachments),
    oldest first. `attachments` is a list of dicts (see
    validate_images()) or None."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT
                    role,
                    content,
                    created_at,
                    id,
                    attachments
                FROM messages
                WHERE conversation_id = %s
                ORDER BY created_at, id;
                """,
                (conversation_id,),
            )

            return cursor.fetchall()


def get_message(conversation_id: int, message_id: int):
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT id, role, content
                FROM messages
                WHERE id = %s AND conversation_id = %s;
                """,
                (message_id, conversation_id),
            )

            return cursor.fetchone()


def get_last_message(conversation_id: int):
    """Returns (id, role, content, attachments) for the most recent
    message, or None."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT id, role, content, attachments
                FROM messages
                WHERE conversation_id = %s
                ORDER BY created_at DESC, id DESC
                LIMIT 1;
                """,
                (conversation_id,),
            )

            return cursor.fetchone()


def delete_messages_from(conversation_id: int, message_id: int):
    """Deletes the message with this id and every message after it in
    the conversation — used for edit (re-ask) and regenerate."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                DELETE FROM messages
                WHERE conversation_id = %s AND id >= %s;
                """,
                (conversation_id, message_id),
            )


# --- Generated files / canvas helpers ---------------------------------
# Backs the generate_document / generate_image / write_code tools (see
# the "Generation" section below) and the design/canvas side panel.

def save_generated_file(
    user_id: int,
    conversation_id: int | None,
    title: str,
    kind: str,
    mime_type: str,
    data: bytes,
) -> int:
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO generated_files (
                    user_id, conversation_id, title, kind, mime_type, data
                )
                VALUES (%s, %s, %s, %s, %s, %s)
                RETURNING id;
                """,
                (user_id, conversation_id, title, kind, mime_type, data),
            )

            return cursor.fetchone()[0]


def get_generated_file(file_id: int, user_id: int):
    """Returns (title, kind, mime_type, data) or None."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT title, kind, mime_type, data
                FROM generated_files
                WHERE id = %s AND user_id = %s;
                """,
                (file_id, user_id),
            )

            return cursor.fetchone()


def upsert_canvas_artifact(
    conversation_id: int,
    title: str,
    kind: str,
    content: str | None,
    language: str | None = None,
    generated_file_id: int | None = None,
):
    """Creates or overwrites the conversation's single canvas
    artifact — the canvas holds one "current working artifact" at a
    time (see database/schema.sql), so a follow-up generation call
    updates it in place rather than creating a new one."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO canvas_artifacts (
                    conversation_id, title, kind, language, content,
                    generated_file_id, updated_at
                )
                VALUES (%s, %s, %s, %s, %s, %s, CURRENT_TIMESTAMP)
                ON CONFLICT (conversation_id) DO UPDATE SET
                    title = EXCLUDED.title,
                    kind = EXCLUDED.kind,
                    language = EXCLUDED.language,
                    content = EXCLUDED.content,
                    generated_file_id = EXCLUDED.generated_file_id,
                    updated_at = CURRENT_TIMESTAMP
                RETURNING id;
                """,
                (
                    conversation_id,
                    title,
                    kind,
                    language,
                    content,
                    generated_file_id,
                ),
            )

            return cursor.fetchone()[0]


def get_canvas_artifact(conversation_id: int):
    """Returns (id, title, kind, language, content, generated_file_id,
    updated_at) for the conversation's canvas artifact, or None."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT
                    id, title, kind, language, content,
                    generated_file_id, updated_at
                FROM canvas_artifacts
                WHERE conversation_id = %s;
                """,
                (conversation_id,),
            )

            return cursor.fetchone()


def serialize_datetime(value: datetime) -> str:
    return value.isoformat()


def build_conversation_label(
    chat_number: int,
    created_at: datetime,
) -> str:
    local_time = created_at.astimezone()

    formatted_time = local_time.strftime("%m/%d/%Y %I:%M %p")

    formatted_time = formatted_time.replace(
        " 0",
        " ",
    ).lower()

    return f"Chat {chat_number} - {formatted_time}"


# --- Model routing -----------------------------------------------------
# Heuristic router: short, simple-looking messages go to the fast/free
# model; anything that looks like it needs real reasoning (long messages,
# code, multi-part questions, explicit "explain/analyze/write" asks) goes
# to the smarter model. This is intentionally simple and cheap to run —
# swap in a small classifier model later if you want more accuracy.
_COMPLEXITY_KEYWORDS = (
    "explain", "analyze", "analyse", "write", "code", "debug",
    "compare", "design", "architecture", "plan", "summarize",
    "summarise", "translate", "why", "how does", "step by step",
    "essay", "report", "proof", "solve", "search", "look up",
    "calculate", "compute",
)

# Correctness fix, not a complexity tweak: FAST_MODEL is never in
# TOOL_CAPABLE_MODELS below (see that comment), so a short message like
# "Generate the pdf containing this" or "make that into slides" — too
# short/plain to trip any _COMPLEXITY_KEYWORDS above — used to route to
# FAST_MODEL, which then had literally no tools available (not
# create_document, not generate_image, none of them) and fell back to
# its own base-training refusal ("I can't generate files...") instead
# of ever reaching a model that actually has the tool. These keywords
# force that class of request onto a tool-capable model regardless of
# length. English file-format tokens (pdf/docx/pptx) are kept
# unmodified across languages since they're commonly used as loanwords
# even in non-English messages; the verbs and the higher-level nouns
# (word doc, presentation, slides, image, ...) are translated for the
# 11 languages EASTA_SYSTEM_PROMPT below claims fluency in.
_GENERATION_KEYWORDS = (
    # verbs: generate/create/make/turn-into/convert
    "generate", "create", "make", "produce", "turn this into",
    "turn that into", "convert this", "convert that",
    "générer", "créer", "crée", "fais",
    "generar", "crear", "crea", "haz",
    "erstelle", "erstellen", "generiere", "generieren",
    "生成", "创建", "制作",
    "создай", "создать", "сгенерируй", "сгенерировать",
    "gerar", "criar", "crie", "faça",
    "作成", "作って",
    "생성", "만들어", "작성",
    "genera", "generare", "creare",
    "genereer", "maak", "creëer",
    # file-format nouns (language-agnostic loanwords)
    "pdf", "docx", "pptx",
    # higher-level nouns: word doc / presentation / slides / image
    "word doc", "powerpoint", "presentation", "slides", "slideshow",
    "image", "picture", "photo", "logo",
    "document word", "présentation", "diapositives",
    "documento word", "presentación", "diapositivas", "imagen",
    "word-dokument", "präsentation", "folien", "bild",
    "word文档", "演示文稿", "幻灯片", "图片", "图像",
    "документ word", "презентация", "слайды", "изображение", "картинка",
    "documento word", "apresentação", "imagem",
    "word文書", "プレゼンテーション", "スライド", "パワーポイント", "画像",
    "워드 문서", "프레젠테이션", "슬라이드", "파워포인트", "이미지",
    "documento word", "presentazione", "diapositive", "immagine",
    "word-document", "presentatie", "dia's", "afbeelding",
)


def pick_model(message: str) -> str:
    text = message.lower()
    word_count = len(text.split())

    looks_complex = (
        word_count > 40
        or "```" in message
        or any(keyword in text for keyword in _COMPLEXITY_KEYWORDS)
        or any(keyword in text for keyword in _GENERATION_KEYWORDS)
    )

    return SMART_MODEL if looks_complex else FAST_MODEL


def model_chain_for(message: str) -> list[str]:
    """Ordered list of models to try: chosen model, then the other
    primary model, then the fallback — so a rate-limit or outage on one
    provider doesn't take down the whole chat."""
    primary = pick_model(message)
    secondary = FAST_MODEL if primary == SMART_MODEL else SMART_MODEL

    chain = [primary, secondary, FALLBACK_MODEL]

    # de-duplicate while preserving order
    seen = set()
    ordered = []
    for model in chain:
        if model not in seen:
            seen.add(model)
            ordered.append(model)

    return ordered


def model_chain_for_images() -> list[str]:
    """Model chain for turns with an image attachment — restricted to
    models confirmed to accept image input, since EASTA_FAST_MODEL and
    EASTA_FALLBACK_MODEL aren't guaranteed to support vision."""
    chain = [VISION_MODEL, VISION_FALLBACK_MODEL]

    seen = set()
    ordered = []
    for model in chain:
        if model not in seen:
            seen.add(model)
            ordered.append(model)

    return ordered


# --- RAG: grounding on the user's own documents -----------------------------
# Keyword search over Postgres tsvector. No embeddings, no extra services,
# works on a stock managed Postgres instance. See documents table in
# database/schema.sql. Swap for pgvector + real embeddings later if you
# need semantic (not just keyword) matching.

def retrieve_relevant_chunks(query: str, user_id: int, limit: int = 3):
    """Phase 22: searches the user's personal documents AND every
    organization they belong to's shared documents in one query (the
    org_id subquery), not just their own -- someone in two orgs gets
    grounding from both, no special-casing needed for "how many orgs"."""
    query = query.strip()

    if not query:
        return []

    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            try:
                cursor.execute(
                    """
                    SELECT
                        title,
                        content,
                        ts_rank_cd(
                            search_vector,
                            websearch_to_tsquery('english', %s)
                        ) AS rank
                    FROM documents
                    WHERE
                        (
                            user_id = %s
                            OR org_id IN (
                                SELECT org_id FROM organization_members
                                WHERE user_id = %s
                            )
                        )
                        AND search_vector @@
                            websearch_to_tsquery('english', %s)
                    ORDER BY rank DESC
                    LIMIT %s;
                    """,
                    (query, user_id, user_id, query, limit),
                )

                return cursor.fetchall()

            except psycopg.Error:
                # websearch_to_tsquery can reject odd input (e.g. bare
                # punctuation); grounding is best-effort, so degrade
                # gracefully instead of failing the whole chat turn.
                connection.rollback()
                return []


def build_rag_system_message(query: str, user_id: int):
    if not ENABLE_RAG:
        return None

    chunks = retrieve_relevant_chunks(query, user_id)

    if not chunks:
        return None

    sections = []

    for title, content, _rank in chunks:
        trimmed = content.strip()[:1200]
        sections.append(f"### {title}\n{trimmed}")

    context_text = "\n\n".join(sections)

    return {
        "role": "system",
        "content": (
            "Grounding context (from the user's own knowledge base, "
            "most relevant first). Use this if it helps answer the "
            "question; say so if it doesn't cover the question:\n\n"
            f"{context_text}"
        ),
    }


# --- Tool use: web search & code execution ----------------------------------

TOOLS = [
    {
        "type": "function",
        "function": {
            "name": "web_search",
            "description": (
                "Search the public web for current or specific "
                "information (news, prices, facts you're not sure "
                "of, anything that could have changed recently)."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "query": {
                        "type": "string",
                        "description": "The search query.",
                    }
                },
                "required": ["query"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "execute_python",
            "description": (
                "Run a short, self-contained Python snippet for exact "
                "calculations or data processing. Use print() to "
                "produce output you want to see. No filesystem, "
                "network, or OS access is available."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "code": {
                        "type": "string",
                        "description": "The Python code to execute.",
                    }
                },
                "required": ["code"],
            },
        },
    },
]


def _resolve_duckduckgo_url(href: str) -> str:
    """DuckDuckGo's HTML result links are redirects through
    duckduckgo.com/l/?uddg=<url-encoded target> rather than direct
    links — unwrap them so citations and fetch_page_text() point at
    the actual source instead of a redirect that fails to fetch."""
    if not href:
        return href

    if href.startswith("//"):
        href = "https:" + href

    parsed = urllib.parse.urlparse(href)

    if parsed.netloc.endswith("duckduckgo.com") and parsed.path == "/l/":
        target = urllib.parse.parse_qs(parsed.query).get("uddg")
        if target:
            return target[0]

    return href


def _duckduckgo_search(query: str, max_results: int = 5) -> list[dict]:
    """Raw DuckDuckGo HTML-scrape search, returning structured
    {title, url, snippet} dicts. Shared by the lightweight web_search
    tool and the research-mode pipeline below.
    ⚠️ Prototype-grade: an HTML scrape is fragile (breaks on markup
    changes, gets rate-limited under load) — swap for a real search API
    (Tavily, Serper, Bing, or similar) before depending on this for
    real traffic. This is the one function to change to do that."""
    query = (query or "").strip()

    if not query:
        return []

    try:
        response = requests.get(
            "https://html.duckduckgo.com/html/",
            params={"q": query},
            headers={
                "User-Agent": "Mozilla/5.0 (compatible; EASTA/1.0)"
            },
            timeout=8,
        )
        response.raise_for_status()

    except requests.RequestException:
        return []

    soup = BeautifulSoup(response.text, "html.parser")
    results = []

    for result in soup.select(".result")[:max_results]:
        title_el = result.select_one(".result__a")

        if not title_el:
            continue

        snippet_el = result.select_one(".result__snippet")

        results.append({
            "title": title_el.get_text(strip=True),
            "url": _resolve_duckduckgo_url(title_el.get("href", "")),
            "snippet": snippet_el.get_text(strip=True) if snippet_el else "",
        })

    return results


def tool_web_search(query: str, max_results: int = 5) -> str:
    query = (query or "").strip()

    if not query:
        return "No search query was provided."

    results = _duckduckgo_search(query, max_results)

    if not results:
        return f"No web results found for '{query}'."

    formatted = [
        f"- {r['title']}\n  {r['url']}\n  {r['snippet']}" for r in results
    ]

    return f"Web search results for '{query}':\n\n" + "\n\n".join(formatted)


# NOTE ON SAFETY: this is a naive blocklist + subprocess sandbox meant
# for a single-tenant prototype only. It is NOT secure isolation for
# untrusted, multi-tenant production use — a determined user can likely
# work around the blocklist. Before shipping this to real users, replace
# the body of this function with a call to a real sandboxing service
# (e.g. E2B, Modal Sandboxes, a locked-down Docker/gVisor container with
# no network and a resource-limited user).
_BLOCKED_CODE_PATTERNS = (
    "import os", "import sys", "import subprocess", "import socket",
    "import shutil", "__import__", "open(", "eval(", "exec(",
    "importlib", "ctypes", "import requests", "urllib",
)


def tool_execute_python(code: str, timeout_seconds: int = 8) -> str:
    code = (code or "").strip()

    if not code:
        return "No code was provided."

    if len(code) > 6000:
        return "Code is too long (6000 character limit)."

    lowered = code.lower()

    for pattern in _BLOCKED_CODE_PATTERNS:
        if pattern in lowered:
            return (
                f"Execution blocked: '{pattern.strip()}' is not "
                "allowed in this prototype sandbox. Only plain "
                "computational Python (math, strings, loops, data "
                "structures) is supported."
            )

    with tempfile.TemporaryDirectory() as tmp_dir:
        script_path = os.path.join(tmp_dir, "snippet.py")

        with open(script_path, "w", encoding="utf-8") as file:
            file.write(code)

        try:
            completed = subprocess.run(
                ["python3", "-I", "-S", script_path],
                cwd=tmp_dir,
                capture_output=True,
                text=True,
                timeout=timeout_seconds,
            )

        except subprocess.TimeoutExpired:
            return f"Execution timed out after {timeout_seconds}s."

        except OSError as error:
            return f"Execution failed to start: {error}"

    output = (completed.stdout or "")[:4000]
    error_output = (completed.stderr or "")[:2000]

    parts = []

    if output:
        parts.append(f"stdout:\n{output}")

    if error_output:
        parts.append(f"stderr:\n{error_output}")

    if not parts:
        parts.append("(no output)")

    return "\n\n".join(parts)


def execute_tool(
    name: str,
    args: dict,
    user_id: int | None = None,
    conversation_id: int | None = None,
) -> tuple[str, dict | None]:
    """Returns (result_text, ui_event). result_text is fed back to the
    model as the tool's output; ui_event (or None) is an extra event
    surfaced to the chat UI -- either a canvas-panel update
    ({"type": "canvas", ...}) or a chat file card
    ({"type": "file_card", ...}) -- see the "Generation" section
    below."""
    try:
        if name == "web_search":
            return tool_web_search(args.get("query", "")), None

        if name == "execute_python":
            return tool_execute_python(args.get("code", "")), None

        if name == "generate_document" and ENABLE_GENERATION:
            return tool_generate_document(args, user_id, conversation_id)

        if name == "generate_image" and ENABLE_GENERATION:
            return tool_generate_image(args, user_id, conversation_id)

        if name == "write_code" and ENABLE_GENERATION:
            return tool_write_code(args, user_id, conversation_id)

        if name == "create_document" and ENABLE_GENERATION:
            return tool_create_document(args, user_id, conversation_id)

        return f"Unknown tool '{name}'.", None

    except Exception as error:  # noqa: BLE001 - tool errors are data
        return f"Tool '{name}' failed: {error}", None


# --- Research mode: multi-query, full-page-read web research ----------------
# A deeper alternative to the single-query web_search tool above: breaks the
# user's question into a few angles, searches each, fetches and reads the
# actual top pages (not just snippets), and hands the model a numbered
# source list to cite from. Still built on the DuckDuckGo HTML scrape (see
# _duckduckgo_search's docstring) — swap that one function for a real
# search API first if you depend on this for real traffic.

MAX_RESEARCH_QUERIES = 3
MAX_RESEARCH_SOURCES = 5
MAX_RESEARCH_PAGE_CHARS = 3000
RESEARCH_FETCH_TIMEOUT = 6


def generate_research_queries(
    topic: str,
    user_id: int | None = None,
    conversation_id: int | None = None,
) -> list[str]:
    """Asks the smart model to break a topic into a few focused search
    queries covering it from different angles. Falls back to the
    original topic verbatim if the call fails or returns nothing
    usable — research mode should degrade, not break, the turn."""
    try:
        response = client.chat.completions.create(
            model=SMART_MODEL,
            messages=[
                {
                    "role": "system",
                    "content": (
                        "Break the user's question into 2-3 focused web "
                        "search queries that together cover it from "
                        "different angles (e.g. background/definition, "
                        "recent developments, specific facts/numbers). "
                        "Reply with ONLY a JSON array of strings, no "
                        "other text."
                    ),
                },
                {"role": "user", "content": topic},
            ],
            max_tokens=200,
        )

        if response.usage and user_id is not None:
            try:
                log_usage(
                    user_id,
                    conversation_id,
                    SMART_MODEL,
                    response.usage.prompt_tokens or 0,
                    response.usage.completion_tokens or 0,
                )
            except Exception as usage_error:  # noqa: BLE001 - best-effort
                print(
                    "EASTA: research query-generation usage logging "
                    f"failed: {usage_error!r}"
                )

        raw = response.choices[0].message.content or "[]"
        queries = json.loads(raw)

        if isinstance(queries, list):
            cleaned = [
                str(q).strip() for q in queries if str(q or "").strip()
            ]
            if cleaned:
                return cleaned[:MAX_RESEARCH_QUERIES]

    except Exception as error:  # noqa: BLE001 - best-effort, has a fallback
        report_error("research query generation failed", error)

    return [topic]


def fetch_page_text(
    url: str,
    max_chars: int = MAX_RESEARCH_PAGE_CHARS,
):
    """Fetches a URL and extracts its main readable text for the
    research pipeline (strips script/style/nav/etc, collapses
    whitespace). Returns None if the page can't be fetched or isn't
    HTML — the caller falls back to the search snippet in that case."""
    try:
        response = requests.get(
            url,
            headers={
                "User-Agent": "Mozilla/5.0 (compatible; EASTA/1.0)"
            },
            timeout=RESEARCH_FETCH_TIMEOUT,
        )
        response.raise_for_status()

    except requests.RequestException:
        return None

    if "html" not in response.headers.get("content-type", ""):
        return None

    soup = BeautifulSoup(response.text, "html.parser")

    for tag in soup(
        ["script", "style", "nav", "header", "footer", "aside", "noscript"]
    ):
        tag.decompose()

    text = " ".join(soup.get_text(separator=" ").split())

    return text[:max_chars] if text else None


def run_research(
    topic: str,
    user_id: int | None = None,
    conversation_id: int | None = None,
):
    """Generator driving research mode. Yields ("meta", dict) progress
    events to surface in the UI while queries run and pages are read,
    then a final ("result", (research_message, sources)) tuple —
    research_message is a system-message dict to inject into the LLM
    context (or None if nothing usable was found), and sources is a
    list of {index, title, url} for the UI's source list under the
    reply. `user_id`/`conversation_id` are only used to attribute the
    query-generation LLM call's cost in usage_logs."""
    queries = generate_research_queries(topic, user_id, conversation_id)

    yield ("meta", {"type": "research_queries", "queries": queries})

    with concurrent.futures.ThreadPoolExecutor(
        max_workers=len(queries)
    ) as executor:
        query_results = executor.map(
            lambda q: _duckduckgo_search(q, max_results=4), queries
        )

    seen_urls = set()
    candidates = []

    for results in query_results:
        for result in results:
            url = result["url"]
            if not url or url in seen_urls:
                continue
            seen_urls.add(url)
            candidates.append(result)

    top_candidates = candidates[:MAX_RESEARCH_SOURCES]

    if not top_candidates:
        yield ("result", (None, []))
        return

    yield ("meta", {
        "type": "research_reading",
        "count": len(top_candidates),
    })

    with concurrent.futures.ThreadPoolExecutor(
        max_workers=len(top_candidates)
    ) as executor:
        future_to_url = {
            executor.submit(fetch_page_text, c["url"]): c["url"]
            for c in top_candidates
        }
        page_text_by_url = {}
        for future in concurrent.futures.as_completed(future_to_url):
            url = future_to_url[future]
            try:
                page_text_by_url[url] = future.result()
            except Exception:  # noqa: BLE001 - fall back to the snippet
                page_text_by_url[url] = None

    sources = []
    context_sections = []

    for candidate in top_candidates:
        body = page_text_by_url.get(candidate["url"]) or candidate["snippet"]

        if not body:
            continue

        index = len(sources) + 1

        sources.append({
            "index": index,
            "title": candidate["title"] or candidate["url"],
            "url": candidate["url"],
        })

        context_sections.append(
            f"[{index}] {candidate['title']}\n{candidate['url']}\n{body}"
        )

    if not context_sections:
        yield ("result", (None, []))
        return

    research_message = {
        "role": "system",
        "content": (
            "Research context — sources gathered for this question, "
            "numbered for citation. Cite them inline like [1], [2] at "
            "the point in your answer that draws on them, and only "
            "cite sources you actually used. If the sources don't "
            "fully answer the question, say so.\n\n"
            + "\n\n".join(context_sections)
        ),
    }

    yield ("result", (research_message, sources))


# --- Generation: documents, images, code, and the canvas panel --------------
# Three tools (generate_document, generate_image, write_code) let the model
# produce a downloadable file and/or push content to the conversation's
# canvas — a single "current working artifact" the side panel displays,
# which a follow-up instruction updates in place (see
# upsert_canvas_artifact()) instead of the model re-pasting the whole
# thing into the chat.

_INLINE_MARKDOWN_RE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


def _markdown_inline_to_reportlab(text: str) -> str:
    """Converts a small subset of inline Markdown (bold/italic/inline
    code) to reportlab's mini-HTML Paragraph markup. Not a full
    Markdown parser — see render_markdown_to_pdf()'s docstring."""
    text = (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"(?<!\*)\*(?!\*)(.+?)\*(?!\*)", r"<i>\1</i>", text)
    text = re.sub(r"`(.+?)`", r"<font face='Courier'>\1</font>", text)
    return text


# --- Full Markdown -> PDF/DOCX, sharing one parse (Phase 15) ---------------
# render_markdown_to_pdf()/render_markdown_to_docx() below used to hand-map
# a narrow Markdown subset directly to reportlab Flowables / python-docx
# calls line-by-line (no tables, nested lists, or images -- see git history
# if you need the old version back). Both now go through one real Markdown
# parser (mistune) into an HTML tree once, and drive their own renderer
# from that same tree, so PDF and DOCX output can't drift apart the way two
# independent hand-rolled walkers eventually do. render_structured_pdf() /
# render_structured_docx() below (the separate create_document tool, a
# structured-JSON-sections content model rather than raw Markdown) are
# untouched and still use reportlab / _markdown_inline_to_reportlab() /
# _add_markdown_runs() directly -- deliberately a different tool, not
# consolidated here.
#
# PDF renders via xhtml2pdf (HTML+CSS -> PDF): chosen over WeasyPrint
# (which the underlying "Suggested next steps" note originally pointed at)
# because WeasyPrint requires native Pango/cairo/GTK libraries that not
# every deployment target has preinstalled (confirmed: it fails to import
# with a missing-libgobject error on a plain pip install on Windows,
# without the GTK3 runtime separately installed) -- xhtml2pdf is pure
# Python, works identically everywhere reportlab already did, and still
# gets tables/nested lists/images "for free" from HTML+CSS the same way.
# If your deployment target can guarantee the native libs (most Linux
# server images can, via apt), swapping the PDF half for WeasyPrint is a
# reasonable follow-up -- the shared markdown_to_soup() parse and the DOCX
# renderer don't need to change either way.
_MARKDOWN_TO_HTML = mistune.create_markdown(
    escape=False,
    plugins=["table", "strikethrough", "task_lists", "url"],
)


def markdown_to_soup(markdown_body: str) -> BeautifulSoup:
    """The single parse both render_markdown_to_pdf() and
    render_markdown_to_docx() are driven from. DOCX walks this tree as
    it comes back; the PDF path takes a further pass (see
    _prepare_pdf_html()) for a couple of xhtml2pdf-specific rendering
    gaps that don't apply to python-docx."""
    parsed_html = _MARKDOWN_TO_HTML(markdown_body or "")
    return BeautifulSoup(parsed_html, "html.parser")


def _body_starts_with_heading(soup: BeautifulSoup) -> bool:
    first_tag = soup.find(True)
    return first_tag is not None and first_tag.name in ("h1", "h2", "h3")


_LONG_TOKEN_RE = re.compile(r"\S{60,}")


def _break_long_tokens(text: str, breaker: str, chunk_size: int = 60) -> str:
    """xhtml2pdf implements neither `word-break` nor `overflow-wrap` --
    confirmed by testing: a single unbroken run of 60+ non-space
    characters (a long URL, hash, minified line, ...) doesn't get
    clipped or raise an error, it silently renders past the page's
    right margin. Inserts a real break every chunk_size characters into
    any such run so it wraps like normal text does; short tokens and
    ordinary prose are left untouched. `breaker` is "\\n" inside a <pre>
    (where whitespace is preserved by `white-space: pre-wrap`) and a
    plain " " everywhere else (a literal "\\n" there would just get
    collapsed back to nothing by normal HTML whitespace handling and
    never actually break the line)."""

    def repl(match):
        token = match.group(0)
        return breaker.join(
            token[i : i + chunk_size] for i in range(0, len(token), chunk_size)
        )

    return _LONG_TOKEN_RE.sub(repl, text)


_PDF_BULLET_CHARS = ["•", "◦", "▪"]


def _ul_nesting_depth(tag: Tag) -> int:
    return sum(
        1 for parent in tag.parents if isinstance(parent, Tag) and parent.name == "ul"
    )


def _prepare_pdf_html(soup: BeautifulSoup) -> str:
    """Two xhtml2pdf-specific fixups, applied to a fresh copy of the
    tree so they never leak into the shared soup the DOCX path also
    reads:
    1. Long-token wrapping (_break_long_tokens), everywhere.
    2. Manual <ul> bullet markers: xhtml2pdf's native disc marker
       doesn't survive registering a custom @font-face body font --
       confirmed via pypdf text extraction, it comes back as a
       replacement-character glyph instead of a bullet. Rendered as a
       literal, nesting-depth-aware character instead (with
       `list-style-type: none` in the CSS below to suppress the native
       one) -- that's just ordinary text in the already-embedded Inter
       font. <ol> numbering is unaffected by this and needs no
       workaround; it already renders correctly."""
    soup = BeautifulSoup(str(soup), "html.parser")

    for pre in soup.find_all("pre"):
        for text_node in pre.find_all(string=True):
            wrapped = _break_long_tokens(str(text_node), breaker="\n")
            if wrapped != str(text_node):
                text_node.replace_with(wrapped)

    for text_node in soup.find_all(string=True):
        if text_node.find_parent("pre") is not None:
            continue
        wrapped = _break_long_tokens(str(text_node), breaker=" ")
        if wrapped != str(text_node):
            text_node.replace_with(wrapped)

    for ul in soup.find_all("ul"):
        marker = _PDF_BULLET_CHARS[
            min(_ul_nesting_depth(ul), len(_PDF_BULLET_CHARS) - 1)
        ]
        for li in ul.find_all("li", recursive=False):
            li.insert(0, NavigableString(f"{marker}  "))

    return str(soup)


# Bundled here (backend/assets/fonts/) rather than relying on default PDF
# fonts, which reportlab/xhtml2pdf can't reach Google Fonts for at render
# time the way a browser's @import can -- same Inter/Fraunces families as
# frontend/static/styles.css, downloaded as static-weight .ttf files.
DOCUMENT_FONTS_DIR = os.path.join(os.path.dirname(__file__), "assets", "fonts")


def _pdf_font_face_css() -> str:
    faces = [
        ("Inter", "normal", "normal", "Inter-Regular.ttf"),
        ("Inter", "bold", "normal", "Inter-Bold.ttf"),
        ("Inter", "normal", "italic", "Inter-Italic.ttf"),
        ("Inter-Bold", "normal", "normal", "Inter-Bold.ttf"),
        ("Inter-SemiBold", "normal", "normal", "Inter-SemiBold.ttf"),
        ("Fraunces-Medium", "normal", "normal", "Fraunces-Medium.ttf"),
        ("Fraunces-SemiBold", "normal", "normal", "Fraunces-SemiBold.ttf"),
    ]
    rules = [
        f'@font-face {{ font-family: "{family}"; '
        f'src: url("{DOCUMENT_FONTS_DIR}/{filename}"); '
        f"font-weight: {weight}; font-style: {style}; }}"
        for family, weight, style, filename in faces
    ]
    return "\n".join(rules)


# Terracotta brand accent + neutrals, matching frontend/static/styles.css
# (--color-accent / --color-accent-hover / --color-text-muted /
# --color-border / --color-bg-soft) -- same values _ACCENT_HEX and friends
# below use for the structured-document renderers, duplicated here as
# plain hex since this runs earlier in the file and CSS wants raw strings
# rather than reportlab/docx color objects anyway.
_PDF_DOCUMENT_CSS = """
@page {
    size: letter;
    margin: 1in 0.9in 0.85in 0.9in;
    @frame footer_frame {
        -pdf-frame-content: footer_content;
        left: 0.9in; right: 0.9in; bottom: 0.35in; height: 0.35in;
    }
}
body {
    font-family: "Inter";
    font-size: 10.5pt;
    line-height: 1.5;
    color: #2A2622;
}
#footer_content {
    text-align: center;
    font-size: 8.5pt;
    color: #6F6558;
}
h1, h2, h3 {
    color: #2A2622;
    margin-top: 18pt;
    margin-bottom: 6pt;
}
h1 { font-family: "Fraunces-SemiBold"; font-size: 20pt; }
h2 { font-family: "Fraunces-Medium"; font-size: 15pt; }
h3 { font-family: "Fraunces-Medium"; font-size: 12.5pt; }
p { margin: 0 0 8pt 0; }
ul, ol { margin: 0 0 8pt 0; padding-left: 18pt; }
ul { list-style-type: none; }
li { margin-bottom: 3pt; }
table {
    width: 100%;
    margin: 4pt 0 10pt 0;
    font-size: 9.5pt;
}
th, td {
    border: 0.75pt solid #D8CBAF;
    padding: 5pt 7pt;
    text-align: left;
}
th {
    background-color: #F1EADC;
    font-family: "Inter-SemiBold";
    color: #AD5731;
}
pre {
    background-color: #F1EADC;
    border: 0.75pt solid #D8CBAF;
    padding: 8pt;
    margin: 0 0 10pt 0;
    font-family: "Courier";
    font-size: 8.5pt;
    white-space: pre-wrap;
}
code {
    font-family: "Courier";
    font-size: 9pt;
    background-color: #F1EADC;
}
blockquote {
    margin: 0 0 8pt 0;
    padding-left: 10pt;
    border-left: 2pt solid #C4693E;
    color: #6F6558;
}
a { color: #AD5731; }
strong { font-family: "Inter-Bold"; }
em { font-family: "Inter"; font-style: italic; }
img { max-width: 100%; }
.title-page {
    page-break-after: always;
    text-align: center;
    padding-top: 2.6in;
}
.title-page h1 {
    font-family: "Fraunces-SemiBold";
    font-size: 30pt;
    margin-bottom: 10pt;
}
.title-page .subtitle {
    font-family: "Inter";
    font-size: 11pt;
    color: #6F6558;
}
"""

# A ~2-page threshold, estimated from character count rather than an
# actual page-layout pass (nothing renders a preview page count ahead of
# the real xhtml2pdf call) -- generous enough that short documents never
# get an unnecessary title page, conservative enough that most documents
# genuinely running past 2 pages get one.
MIN_CHARS_FOR_DOCUMENT_TITLE_PAGE = 2800


def render_markdown_to_pdf(title: str, markdown_body: str) -> bytes:
    soup = markdown_to_soup(markdown_body)
    safe_title = html.escape(title)
    already_has_heading = _body_starts_with_heading(soup)
    body_html = _prepare_pdf_html(soup)

    needs_title_page = len(markdown_body or "") > MIN_CHARS_FOR_DOCUMENT_TITLE_PAGE

    title_page_html = ""
    heading_html = ""
    if needs_title_page:
        # The title page always shows the title once, on its own page
        # -- the body's own first heading (if any) then opens the next
        # page normally, no dedup needed there.
        title_page_html = (
            f'<div class="title-page"><h1>{safe_title}</h1>'
            f'<div class="subtitle">Generated by EASTA</div></div>'
        )
    elif not already_has_heading:
        # Short document, no title page: inject the title as the
        # visible heading -- but only if the model's own Markdown
        # doesn't already open with one, otherwise this doubles up
        # ("Report" / "Report") right at the top with no separation
        # between them (confirmed reproducible without this check).
        heading_html = f"<h1>{safe_title}</h1>"

    full_html = f"""<html><head><style>
{_pdf_font_face_css()}
{_PDF_DOCUMENT_CSS}
</style></head>
<body>
{title_page_html}
{heading_html}
{body_html}
<div id="footer_content">Page <pdf:pagenumber /> of <pdf:pagecount /></div>
</body></html>"""

    buffer = io.BytesIO()
    result = pisa.CreatePDF(full_html, dest=buffer)
    if result.err:
        raise RuntimeError(f"PDF generation failed with {result.err} error(s)")
    return buffer.getvalue()


def _add_markdown_runs(paragraph, text: str):
    """Adds runs to a python-docx paragraph, applying bold/italic/code
    for the same small subset of inline Markdown as
    _markdown_inline_to_reportlab()."""
    for token in _INLINE_MARKDOWN_RE.split(text):
        if not token:
            continue

        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("*") and token.endswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
        else:
            paragraph.add_run(token)


def _docx_shade_cell(cell, hex_color: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shading)


def _docx_walk_inline(paragraph, node, bold=False, italic=False, code=False, strike=False):
    """Recursively adds runs to a docx paragraph for a parsed HTML
    node's inline content (headings, paragraphs, list items, table
    cells, blockquotes all go through this) -- composes nested
    bold/italic/code/strike correctly (e.g. **_both_** inside `code`)
    since the active styles are threaded down as arguments rather than
    reset at each tag."""
    for child in node.children:
        if isinstance(child, NavigableString):
            text = str(child)
            if not text:
                continue
            run = paragraph.add_run(text)
            run.bold = bold or None
            run.italic = italic or None
            if strike:
                run.font.strike = True
            if code:
                run.font.name = "Courier New"
                run.font.size = Pt(9.5)
            continue

        if not isinstance(child, Tag):
            continue

        tag = child.name
        if tag in ("strong", "b"):
            _docx_walk_inline(paragraph, child, True, italic, code, strike)
        elif tag in ("em", "i"):
            _docx_walk_inline(paragraph, child, bold, True, code, strike)
        elif tag == "code":
            _docx_walk_inline(paragraph, child, bold, italic, True, strike)
        elif tag in ("del", "s"):
            _docx_walk_inline(paragraph, child, bold, italic, code, True)
        elif tag == "br":
            paragraph.add_run().add_break()
        elif tag == "a":
            text = child.get_text()
            href = child.get("href", "")
            label = f"{text} ({href})" if href and href != text else text
            run = paragraph.add_run(label)
            run.bold = bold or None
            run.italic = italic or None
            run.font.color.rgb = DOCX_ACCENT
            run.underline = True
        elif tag in ("ul", "ol"):
            # A nested list inside this <li> -- rendered as its own
            # paragraphs by _docx_render_list's caller, not flattened
            # into this inline run (this exact bug -- nested-list text
            # duplicated into both its own paragraphs and its parent
            # li's paragraph -- was caught and fixed during testing).
            continue
        else:
            _docx_walk_inline(paragraph, child, bold, italic, code, strike)


_DOCX_LIST_STYLES_BULLET = ["List Bullet", "List Bullet 2", "List Bullet 3"]
_DOCX_LIST_STYLES_NUMBER = ["List Number", "List Number 2", "List Number 3"]


def _docx_render_list(document, list_tag, level=0):
    ordered = list_tag.name == "ol"
    styles = _DOCX_LIST_STYLES_NUMBER if ordered else _DOCX_LIST_STYLES_BULLET
    style_name = styles[min(level, len(styles) - 1)]

    for li in list_tag.find_all("li", recursive=False):
        content_source = li.find("p", recursive=False) or li
        paragraph = document.add_paragraph(style=style_name)
        _docx_walk_inline(paragraph, content_source)

        for nested in li.find_all(["ul", "ol"], recursive=False):
            _docx_render_list(document, nested, level=level + 1)


def _docx_render_table(document, table_tag):
    rows = table_tag.find_all("tr")
    if not rows:
        return

    column_count = max(len(row.find_all(["th", "td"])) for row in rows)
    table = document.add_table(rows=0, cols=column_count)
    table.style = "Table Grid"

    for row_index, row in enumerate(rows):
        cells = row.find_all(["th", "td"])
        is_header = row_index == 0 and row.find("th") is not None
        docx_row = table.add_row()

        for col_index in range(column_count):
            docx_cell = docx_row.cells[col_index]
            if col_index >= len(cells):
                continue
            paragraph = docx_cell.paragraphs[0]
            _docx_walk_inline(paragraph, cells[col_index])
            if is_header:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = DOCX_ACCENT_DARK
                _docx_shade_cell(docx_cell, "F1EADC")


def _docx_add_image(document, img_tag):
    src = img_tag.get("src", "")
    if not src:
        return
    try:
        response = requests.get(src, timeout=8, stream=True)
        response.raise_for_status()
        data = response.content[: 8 * 1024 * 1024]
        document.add_picture(io.BytesIO(data), width=Inches(5.5))
    except Exception:
        alt = img_tag.get("alt") or "image"
        placeholder = document.add_paragraph(f"[image: {alt}]")
        for run in placeholder.runs:
            run.italic = True


def _docx_render_block(document, node):
    for child in node.children:
        if isinstance(child, NavigableString):
            continue
        if not isinstance(child, Tag):
            continue

        tag = child.name
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            heading = document.add_heading(level=int(tag[1]))
            _docx_walk_inline(heading, child)
        elif tag == "p":
            images = child.find_all("img")
            if images:
                for img in images:
                    _docx_add_image(document, img)
                if child.get_text().strip():
                    _docx_walk_inline(document.add_paragraph(), child)
            else:
                _docx_walk_inline(document.add_paragraph(), child)
        elif tag in ("ul", "ol"):
            _docx_render_list(document, child, level=0)
        elif tag == "table":
            _docx_render_table(document, child)
        elif tag == "pre":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.2)
            run = paragraph.add_run(child.get_text().rstrip("\n"))
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        elif tag == "blockquote":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.3)
            content_source = child.find("p", recursive=False) or child
            _docx_walk_inline(paragraph, content_source)
            for run in paragraph.runs:
                run.italic = True
                run.font.color.rgb = DOCX_MUTED
        elif tag == "hr":
            document.add_paragraph("—" * 20)
        elif tag == "img":
            _docx_add_image(document, child)
        else:
            _docx_render_block(document, child)


def render_markdown_to_docx(title: str, markdown_body: str) -> bytes:
    soup = markdown_to_soup(markdown_body)

    document = DocxDocument()
    if not _body_starts_with_heading(soup):
        # Same dedup as render_markdown_to_pdf(): don't show the title
        # twice when the model's own Markdown already opens with an H1.
        document.add_heading(title, level=0)

    _docx_render_block(document, soup)

    buffer = io.BytesIO()
    document.save(buffer)

    return buffer.getvalue()


# Composer-facing aspect ratio labels -> OpenRouter Image API's
# `aspect_ratio` values (confirmed via OpenRouter's docs: it accepts
# ratio strings like "1:1"/"16:9"/"9:16"/"4:3"/"3:4"/"auto", and support
# can vary by model -- re-check https://openrouter.ai/docs if you swap
# EASTA_IMAGE_MODEL for something that doesn't take this parameter).
IMAGE_ASPECT_RATIOS = {
    "square": "1:1",
    "portrait": "3:4",
    "landscape": "4:3",
}


def generate_image_bytes(
    prompt: str, aspect_ratio: str | None = None
) -> tuple[bytes, str, float]:
    """Calls OpenRouter's dedicated Image API (POST /api/v1/images --
    distinct from the chat completions endpoint used everywhere else in
    this file). Returns (image_bytes, mime_type, cost_usd). Raises on
    failure; callers turn that into a user-facing tool error message."""
    request_body = {"model": IMAGE_MODEL, "prompt": prompt}

    ratio_value = IMAGE_ASPECT_RATIOS.get(aspect_ratio)
    if ratio_value:
        request_body["aspect_ratio"] = ratio_value

    response = requests.post(
        "https://openrouter.ai/api/v1/images",
        headers={
            "Authorization": f"Bearer {OPENROUTER_API_KEY}",
            "Content-Type": "application/json",
        },
        json=request_body,
        timeout=90,
    )
    response.raise_for_status()

    payload = response.json()
    images = payload.get("data") or []

    if not images:
        raise RuntimeError("The image model returned no image.")

    first = images[0]
    image_bytes = base64.b64decode(first["b64_json"])
    mime_type = first.get("media_type", "image/png")
    cost_usd = float((payload.get("usage") or {}).get("cost", 0) or 0)

    if not cost_usd:
        # The dedicated Image API normally reports a real per-call
        # cost in usage.cost (confirmed against the live endpoint) --
        # this is a rough fallback estimate only, so a usage_logs row
        # still gets a non-zero, non-token-based cost if that field is
        # ever missing. Reconcile against your OpenRouter invoice.
        cost_usd = IMAGE_GENERATION_FALLBACK_COST_USD.get(
            IMAGE_MODEL, DEFAULT_IMAGE_GENERATION_FALLBACK_COST_USD
        )

    return image_bytes, mime_type, cost_usd


def log_direct_cost(
    user_id: int,
    conversation_id: int | None,
    model: str,
    cost_usd: float,
):
    """Logs a usage_logs row from a known dollar cost (e.g. from the
    Image API's response) rather than the token-based pricing table in
    log_usage() below."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO usage_logs (
                    user_id, conversation_id, model,
                    prompt_tokens, completion_tokens, cost_usd
                )
                VALUES (%s, %s, %s, 0, 0, %s);
                """,
                (user_id, conversation_id, model, cost_usd),
            )


# --- Voice: server-side STT/TTS fallbacks ------------------------------
# Only reached when the browser's native SpeechRecognition/speechSynthesis
# APIs aren't available (see ENABLE_SERVER_STT/ENABLE_SERVER_TTS and
# GET /api/features) -- e.g. Firefox has no SpeechRecognition support at
# all, and speechSynthesis voice availability/quality varies a lot by OS.

MAX_TTS_CHARS = 4000


def transcribe_audio_bytes(data: bytes, audio_format: str) -> tuple[str, float]:
    """POSTs audio to OpenRouter's transcription endpoint. Returns
    (transcript_text, cost_usd). Raises on failure."""
    response = requests.post(
        "https://openrouter.ai/api/v1/audio/transcriptions",
        headers={"Authorization": f"Bearer {OPENROUTER_API_KEY}"},
        json={
            "model": STT_MODEL,
            "input_audio": {
                "data": base64.b64encode(data).decode("ascii"),
                "format": audio_format,
            },
        },
        timeout=60,
    )
    response.raise_for_status()

    payload = response.json()
    text = payload.get("text", "")
    cost_usd = float((payload.get("usage") or {}).get("cost", 0) or 0)

    return text, cost_usd


# Unlike the transcription/image APIs, the speech endpoint's response
# is raw audio, not JSON -- OpenRouter doesn't return a per-call dollar
# cost for it the way it does elsewhere. Every other paid call in this
# file logs a real returned cost; TTS instead uses a flat per-character
# estimate (OpenAI's tts-1 list price, ~$15/1M characters, as a
# reasonable reference point) so it still shows up in usage_logs rather
# than going uncounted. Reconcile against your OpenRouter invoice
# periodically, same as the token-based pricing table.
TTS_FALLBACK_COST_USD_PER_1K_CHARS = 0.015


def synthesize_speech_bytes(text: str) -> tuple[bytes, float]:
    """POSTs text to OpenRouter's speech endpoint. Returns (audio_mp3,
    estimated_cost_usd). Raises on failure."""
    response = requests.post(
        "https://openrouter.ai/api/v1/audio/speech",
        headers={
            "Authorization": f"Bearer {OPENROUTER_API_KEY}",
            "Content-Type": "application/json",
        },
        json={
            "model": TTS_MODEL,
            "input": text,
            "voice": TTS_VOICE,
            "response_format": "mp3",
        },
        timeout=60,
    )
    response.raise_for_status()

    cost_usd = (len(text) / 1000) * TTS_FALLBACK_COST_USD_PER_1K_CHARS

    return response.content, cost_usd


GENERATION_TOOLS = [
    {
        "type": "function",
        "function": {
            "name": "generate_document",
            "description": (
                "Generate a downloadable PDF or DOCX document from "
                "Markdown content -- e.g. when the user asks for a "
                "report, a write-up, or something 'as a PDF/Word "
                "file'. Also pushes the document to the canvas side "
                "panel. Call this again with the full updated Markdown "
                "to revise a document already on the canvas -- do not "
                "paste the document text in your reply either way, "
                "just briefly confirm what you made or changed."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "title": {
                        "type": "string",
                        "description": "Short document title, used as the filename and canvas heading.",
                    },
                    "format": {
                        "type": "string",
                        "enum": ["pdf", "docx"],
                    },
                    "markdown": {
                        "type": "string",
                        "description": "The full document body as Markdown (headings, lists, bold/italic, code blocks).",
                    },
                },
                "required": ["title", "format", "markdown"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "generate_image",
            "description": (
                "Generate an image from a text description, shown "
                "inline in your reply with a download button and a "
                "regenerate action -- use only when the user "
                "explicitly asks for an image/picture/logo/"
                "illustration to be created, not for finding existing "
                "images (use web_search for that)."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "prompt": {
                        "type": "string",
                        "description": "A detailed description of the image to generate.",
                    },
                    "aspect_ratio": {
                        "type": "string",
                        "enum": ["square", "portrait", "landscape"],
                        "description": (
                            "Image shape. Default to 'square' unless "
                            "the request clearly calls for something "
                            "else (e.g. a phone wallpaper -> "
                            "'portrait', a banner/wide scene -> "
                            "'landscape')."
                        ),
                    },
                },
                "required": ["prompt"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "write_code",
            "description": (
                "Write or revise a piece of code on the canvas side "
                "panel -- use when the user wants a script/program/"
                "snippet to iterate on, not for a one-line inline "
                "example. Call again with the full updated code to "
                "revise code already on the canvas instead of pasting "
                "the whole thing in your reply -- just briefly confirm "
                "what changed."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "title": {
                        "type": "string",
                        "description": "Short title for the canvas heading, e.g. 'Fibonacci script'.",
                    },
                    "language": {
                        "type": "string",
                        "description": "Language for syntax highlighting, e.g. 'python', 'javascript'.",
                    },
                    "code": {
                        "type": "string",
                        "description": "The full source code.",
                    },
                },
                "required": ["title", "language", "code"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "create_document",
            "description": (
                "Create a polished, downloadable PDF, DOCX, or PPTX "
                "file from structured content (sections with headings, "
                "paragraphs, bullet lists, and/or a table) -- use this "
                "instead of generate_document when the user wants a "
                "finished, well-typeset deliverable (a report, a slide "
                "deck, a document with real tables) rather than "
                "something to keep iterating on in the canvas. Shows "
                "up in the chat as a downloadable file card. Do not "
                "paste the content in your reply -- just briefly "
                "confirm what you made."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "format": {
                        "type": "string",
                        "enum": ["pdf", "docx", "pptx"],
                        "description": (
                            "pptx produces a title slide plus one "
                            "slide per section."
                        ),
                    },
                    "title": {
                        "type": "string",
                        "description": (
                            "Document/deck title, used as the "
                            "filename and title page/slide."
                        ),
                    },
                    "sections": {
                        "type": "array",
                        "description": (
                            "Content sections, in order. For pptx "
                            "each section becomes one slide -- keep "
                            "bullets short and few per section for "
                            "slides."
                        ),
                        "items": {
                            "type": "object",
                            "properties": {
                                "heading": {
                                    "type": "string",
                                    "description": "Section heading / slide title.",
                                },
                                "paragraphs": {
                                    "type": "array",
                                    "items": {"type": "string"},
                                    "description": (
                                        "Body paragraphs (omit for a "
                                        "bullet-only or table-only "
                                        "section)."
                                    ),
                                },
                                "bullets": {
                                    "type": "array",
                                    "items": {"type": "string"},
                                    "description": "Bullet points.",
                                },
                                "table": {
                                    "type": "object",
                                    "description": "An optional table for this section.",
                                    "properties": {
                                        "headers": {
                                            "type": "array",
                                            "items": {"type": "string"},
                                        },
                                        "rows": {
                                            "type": "array",
                                            "items": {
                                                "type": "array",
                                                "items": {"type": "string"},
                                            },
                                        },
                                    },
                                },
                            },
                            "required": ["heading"],
                        },
                    },
                },
                "required": ["format", "title", "sections"],
            },
        },
    },
]


def tools_for_model() -> list[dict]:
    tools = list(TOOLS)

    if ENABLE_GENERATION:
        tools += GENERATION_TOOLS

    return tools


def tool_generate_document(
    args: dict, user_id: int, conversation_id: int
) -> tuple[str, dict | None]:
    title = (args.get("title") or "Document").strip()[:150] or "Document"
    file_format = (args.get("format") or "pdf").strip().lower()
    markdown_body = args.get("markdown") or ""

    if file_format not in ("pdf", "docx"):
        return f"Unsupported format '{file_format}'. Use 'pdf' or 'docx'.", None

    if not markdown_body.strip():
        return "No document content was provided.", None

    if file_format == "pdf":
        file_bytes = render_markdown_to_pdf(title, markdown_body)
        mime_type = "application/pdf"
    else:
        file_bytes = render_markdown_to_docx(title, markdown_body)
        mime_type = DOCX_MIME_TYPE

    file_id = save_generated_file(
        user_id, conversation_id, title, file_format, mime_type, file_bytes
    )
    upsert_canvas_artifact(
        conversation_id,
        title=title,
        kind="document",
        content=markdown_body,
        generated_file_id=file_id,
    )

    result_text = (
        f"Generated a {file_format.upper()} titled \"{title}\" and "
        "placed it on the canvas panel with a download button. Do not "
        "repeat the document's text in your reply -- just briefly "
        "confirm what you made."
    )

    canvas_event = {
        "type": "canvas",
        "title": title,
        "kind": "document",
        "language": None,
        "content": markdown_body,
        "download_url": f"/api/generated/{file_id}/download",
    }

    return result_text, canvas_event


def tool_generate_image(
    args: dict, user_id: int, conversation_id: int
) -> tuple[str, dict | None]:
    prompt = (args.get("prompt") or "").strip()

    if not prompt:
        return "No image prompt was provided.", None

    aspect_ratio = args.get("aspect_ratio")
    if aspect_ratio not in IMAGE_ASPECT_RATIOS:
        aspect_ratio = "square"

    image_bytes, mime_type, cost_usd = generate_image_bytes(prompt, aspect_ratio)

    title = prompt[:150]
    file_id = save_generated_file(
        user_id, conversation_id, title, "image", mime_type, image_bytes
    )

    # Always logged, even if generate_image_bytes() had to fall back to
    # an estimated cost -- every generation call should show up in
    # /usage, not just the ones with a nonzero real cost.
    try:
        log_direct_cost(user_id, conversation_id, IMAGE_MODEL, cost_usd)
    except Exception as error:  # noqa: BLE001 - best-effort
        report_error("image-generation cost logging failed", error)

    result_text = (
        f"Generated an image for: \"{prompt}\" ({aspect_ratio}) and "
        "shown it inline in the chat with a download button and a "
        "regenerate action. Do not describe it in detail unless "
        "asked -- just briefly confirm what you made."
    )

    ui_event = {
        "type": "image_result",
        "id": file_id,
        "prompt": prompt,
        "aspect_ratio": aspect_ratio,
        "download_url": f"/api/generated/{file_id}/download",
    }

    return result_text, ui_event


def tool_write_code(
    args: dict, _user_id: int, conversation_id: int
) -> tuple[str, dict | None]:
    title = (args.get("title") or "Code").strip()[:150] or "Code"
    language = (args.get("language") or "text").strip()[:40] or "text"
    code = args.get("code") or ""

    if not code.strip():
        return "No code was provided.", None

    upsert_canvas_artifact(
        conversation_id,
        title=title,
        kind="code",
        content=code,
        language=language,
    )

    result_text = (
        f"Wrote \"{title}\" ({language}) to the canvas panel. Do not "
        "paste the code in your reply -- just briefly confirm what you "
        "made or changed."
    )

    canvas_event = {
        "type": "canvas",
        "title": title,
        "kind": "code",
        "language": language,
        "content": code,
        "download_url": None,
    }

    return result_text, canvas_event


# --- create_document: styled PDF / DOCX / PPTX from structured content ------
# A more capable sibling of generate_document above: instead of a Markdown
# blob, the model sends structured sections (heading, paragraphs, bullets,
# an optional table), rendered with real typography/native styles per
# format and shown in the chat as a downloadable file card (not pushed to
# the canvas -- this is meant as a finished deliverable, not a draft to
# keep iterating on).

DOCX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
PPTX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument"
    ".presentationml.presentation"
)

# Terracotta brand accent, matching frontend/static/styles.css's
# --color-accent (#C4693E) / --color-accent-hover (#AD5731) / muted text
# (#6F6558). Kept in sync manually -- reportlab, python-docx, and
# python-pptx each have their own incompatible RGB color type.
_ACCENT_HEX = "C4693E"
_ACCENT_DARK_HEX = "AD5731"
_MUTED_HEX = "6F6558"

PDF_ACCENT = reportlab_colors.HexColor(f"#{_ACCENT_HEX}")
PDF_ACCENT_DARK = reportlab_colors.HexColor(f"#{_ACCENT_DARK_HEX}")
PDF_MUTED = reportlab_colors.HexColor(f"#{_MUTED_HEX}")
PDF_BORDER = reportlab_colors.HexColor("#D8CBAF")

DOCX_ACCENT = DocxRGBColor.from_string(_ACCENT_HEX)
DOCX_ACCENT_DARK = DocxRGBColor.from_string(_ACCENT_DARK_HEX)
DOCX_MUTED = DocxRGBColor.from_string(_MUTED_HEX)

PPTX_ACCENT = PptxRGBColor.from_string(_ACCENT_HEX)
PPTX_ACCENT_DARK = PptxRGBColor.from_string(_ACCENT_DARK_HEX)
PPTX_WHITE = PptxRGBColor.from_string("FFFFFF")

# Keeps a single request from producing an oversized file -- sections/
# slides beyond the cap are dropped (truncated=True is reported back to
# the model and the caller), and every text field is clipped.
MAX_DOCUMENT_SECTIONS = 20
MAX_PPTX_SECTIONS = 15
MAX_PARAGRAPHS_PER_SECTION = 6
MAX_BULLETS_PER_SECTION = 10
MAX_TABLE_ROWS = 25
MAX_TABLE_COLUMNS = 8
MAX_HEADING_CHARS = 150
MAX_PARAGRAPH_CHARS = 1000
MAX_BULLET_CHARS = 300
MAX_TABLE_CELL_CHARS = 200


def _clip(text, limit: int) -> str:
    return str(text or "").strip()[:limit]


def sanitize_document_sections(raw_sections, document_format: str):
    """Validates and clamps a create_document tool call's `sections`
    payload to the limits above. Returns (sections, truncated) -- a
    section/list that's too long is cut, not rejected, so the model
    still gets a usable (if smaller) document back."""
    section_limit = (
        MAX_PPTX_SECTIONS if document_format == "pptx" else MAX_DOCUMENT_SECTIONS
    )

    if not isinstance(raw_sections, list):
        return [], False

    truncated = len(raw_sections) > section_limit
    sections = []

    for raw_section in raw_sections[:section_limit]:
        if not isinstance(raw_section, dict):
            continue

        heading = _clip(raw_section.get("heading"), MAX_HEADING_CHARS)

        raw_paragraphs = raw_section.get("paragraphs") or []
        if len(raw_paragraphs) > MAX_PARAGRAPHS_PER_SECTION:
            truncated = True
        paragraphs = [
            _clip(p, MAX_PARAGRAPH_CHARS)
            for p in raw_paragraphs[:MAX_PARAGRAPHS_PER_SECTION]
            if str(p or "").strip()
        ]

        raw_bullets = raw_section.get("bullets") or []
        if len(raw_bullets) > MAX_BULLETS_PER_SECTION:
            truncated = True
        bullets = [
            _clip(b, MAX_BULLET_CHARS)
            for b in raw_bullets[:MAX_BULLETS_PER_SECTION]
            if str(b or "").strip()
        ]

        table = None
        raw_table = raw_section.get("table")

        if isinstance(raw_table, dict):
            raw_headers = raw_table.get("headers") or []
            raw_rows = raw_table.get("rows") or []

            if len(raw_headers) > MAX_TABLE_COLUMNS or len(raw_rows) > MAX_TABLE_ROWS:
                truncated = True

            headers = [
                _clip(h, MAX_TABLE_CELL_CHARS)
                for h in raw_headers[:MAX_TABLE_COLUMNS]
            ]
            rows = [
                [
                    _clip(cell, MAX_TABLE_CELL_CHARS)
                    for cell in row[:MAX_TABLE_COLUMNS]
                ]
                for row in raw_rows[:MAX_TABLE_ROWS]
                if isinstance(row, list)
            ]

            if headers or rows:
                table = {"headers": headers, "rows": rows}

        if not heading and not paragraphs and not bullets and not table:
            continue

        sections.append({
            "heading": heading or "Untitled section",
            "paragraphs": paragraphs,
            "bullets": bullets,
            "table": table,
        })

    return sections, truncated


def _build_pdf_table(table: dict):
    headers = table.get("headers") or []
    rows = table.get("rows") or []

    data = ([headers] if headers else []) + rows

    if not data:
        return Spacer(1, 0)

    pdf_table = Table(data, hAlign="LEFT")

    style_commands = [
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, PDF_BORDER),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]

    if headers:
        style_commands += [
            ("BACKGROUND", (0, 0), (-1, 0), PDF_ACCENT),
            ("TEXTCOLOR", (0, 0), (-1, 0), reportlab_colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ]

    pdf_table.setStyle(TableStyle(style_commands))
    return pdf_table


def render_structured_pdf(title: str, sections: list[dict]) -> bytes:
    """Real typography, not a monospace dump: a title page for longer
    documents, a colored heading hierarchy, and styled tables --
    EASTA's terracotta accent on headings/table headers, Helvetica
    instead of the reportlab default Times-Roman."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=LETTER,
        topMargin=1 * inch,
        bottomMargin=1 * inch,
        leftMargin=1 * inch,
        rightMargin=1 * inch,
        title=title,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "EastaTitle", parent=styles["Title"],
        fontName="Helvetica-Bold", fontSize=26, leading=32,
        textColor=PDF_ACCENT_DARK,
    )
    subtitle_style = ParagraphStyle(
        "EastaSubtitle", parent=styles["Normal"],
        fontName="Helvetica", fontSize=11, leading=15,
        textColor=PDF_MUTED, alignment=1,
    )
    heading_style = ParagraphStyle(
        "EastaHeading1", parent=styles["Heading1"],
        fontName="Helvetica-Bold", fontSize=16, leading=20,
        textColor=PDF_ACCENT, spaceBefore=16, spaceAfter=8,
    )
    body_style = ParagraphStyle(
        "EastaBody", parent=styles["BodyText"],
        fontName="Helvetica", fontSize=10.5, leading=15.5,
        spaceAfter=8,
    )
    bullet_style = ParagraphStyle("EastaBullet", parent=body_style, leftIndent=6)

    story = []
    use_title_page = len(sections) >= 3

    if use_title_page:
        story.append(Spacer(1, 2.2 * inch))
        story.append(Paragraph(_markdown_inline_to_reportlab(title), title_style))
        story.append(Spacer(1, 8))
        story.append(Paragraph(
            f"Generated by EASTA &middot; {datetime.now():%B %d, %Y}",
            subtitle_style,
        ))
        story.append(PageBreak())
    else:
        story.append(Paragraph(_markdown_inline_to_reportlab(title), title_style))
        story.append(Spacer(1, 14))

    for section in sections:
        story.append(Paragraph(
            _markdown_inline_to_reportlab(section.get("heading", "")),
            heading_style,
        ))

        for paragraph_text in section.get("paragraphs") or []:
            story.append(Paragraph(
                _markdown_inline_to_reportlab(paragraph_text), body_style
            ))

        bullets = section.get("bullets") or []
        if bullets:
            story.append(ListFlowable(
                [
                    ListItem(Paragraph(_markdown_inline_to_reportlab(b), bullet_style))
                    for b in bullets
                ],
                bulletType="bullet",
            ))
            story.append(Spacer(1, 6))

        table = section.get("table")
        if table:
            story.append(_build_pdf_table(table))
            story.append(Spacer(1, 10))

    doc.build(story)
    return buffer.getvalue()


def _color_docx_runs(paragraph, rgb: DocxRGBColor):
    for run in paragraph.runs:
        run.font.color.rgb = rgb


def _add_docx_table(document, table: dict):
    headers = table.get("headers") or []
    rows = table.get("rows") or []

    column_count = len(headers) if headers else (len(rows[0]) if rows else 0)
    if column_count == 0:
        return

    total_rows = (1 if headers else 0) + len(rows)
    docx_table = document.add_table(rows=total_rows, cols=column_count)
    docx_table.style = "Table Grid"

    row_index = 0
    if headers:
        for col_index, header_text in enumerate(headers):
            cell = docx_table.cell(0, col_index)
            run = cell.paragraphs[0].add_run(header_text)
            run.bold = True
            run.font.color.rgb = DOCX_ACCENT
        row_index = 1

    for row in rows:
        for col_index in range(column_count):
            value = row[col_index] if col_index < len(row) else ""
            docx_table.cell(row_index, col_index).text = str(value)
        row_index += 1

    document.add_paragraph()


def render_structured_docx(title: str, sections: list[dict]) -> bytes:
    """Uses native Word styles (Title/Heading 1/List Bullet/Table Grid)
    so the output is editable and looks like a real Word document, not
    manually-formatted runs -- with EASTA's terracotta on headings and
    a page break before the body for longer documents."""
    document = DocxDocument()

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)

    title_paragraph = document.add_heading(title, level=0)
    _color_docx_runs(title_paragraph, DOCX_ACCENT_DARK)

    if len(sections) >= 3:
        subtitle = document.add_paragraph(
            f"Generated by EASTA · {datetime.now():%B %d, %Y}"
        )
        subtitle.runs[0].font.size = Pt(10)
        subtitle.runs[0].font.color.rgb = DOCX_MUTED
        document.add_page_break()

    for section in sections:
        heading_paragraph = document.add_heading(
            section.get("heading", ""), level=1
        )
        _color_docx_runs(heading_paragraph, DOCX_ACCENT)

        for paragraph_text in section.get("paragraphs") or []:
            _add_markdown_runs(document.add_paragraph(), paragraph_text)

        for bullet_text in section.get("bullets") or []:
            _add_markdown_runs(
                document.add_paragraph(style="List Bullet"), bullet_text
            )

        table = section.get("table")
        if table:
            _add_docx_table(document, table)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _add_pptx_table(slide, table: dict):
    headers = table.get("headers") or []
    rows = table.get("rows") or []

    column_count = len(headers) if headers else (len(rows[0]) if rows else 0)
    if column_count == 0:
        return

    row_count = min((1 if headers else 0) + len(rows), 9)

    graphic_frame = slide.shapes.add_table(
        row_count, column_count,
        PptxInches(0.6), PptxInches(1.8), PptxInches(9.0), PptxInches(0.5 * row_count),
    )
    pptx_table = graphic_frame.table

    row_index = 0
    if headers:
        for col_index, header_text in enumerate(headers[:column_count]):
            cell = pptx_table.cell(0, col_index)
            cell.text = header_text
            cell.fill.solid()
            cell.fill.fore_color.rgb = PPTX_ACCENT
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.bold = True
                paragraph.font.color.rgb = PPTX_WHITE
        row_index = 1

    for row in rows:
        if row_index >= row_count:
            break
        for col_index in range(column_count):
            value = row[col_index] if col_index < len(row) else ""
            pptx_table.cell(row_index, col_index).text = str(value)
        row_index += 1


def render_structured_pptx(title: str, sections: list[dict]) -> bytes:
    """One slide per section on a simple, consistent title+content
    layout (never walls of text -- bullets/paragraphs are already
    clamped short by sanitize_document_sections), plus a title slide.
    EASTA's terracotta on every slide title."""
    presentation = Presentation()

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = title
    title_paragraph = title_slide.shapes.title.text_frame.paragraphs[0]
    title_paragraph.font.bold = True
    title_paragraph.font.color.rgb = PPTX_ACCENT_DARK

    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = (
            f"Generated by EASTA · {datetime.now():%B %d, %Y}"
        )

    content_layout = presentation.slide_layouts[1]

    for section in sections:
        slide = presentation.slides.add_slide(content_layout)
        slide.shapes.title.text = section.get("heading", "")
        title_run = slide.shapes.title.text_frame.paragraphs[0]
        title_run.font.bold = True
        title_run.font.color.rgb = PPTX_ACCENT

        body_placeholder = None
        for placeholder in slide.placeholders:
            if placeholder.placeholder_format.idx != 0:
                body_placeholder = placeholder
                break

        table = section.get("table")
        bullets = section.get("bullets") or []
        paragraphs = section.get("paragraphs") or []

        if table:
            _add_pptx_table(slide, table)
        elif bullets and body_placeholder:
            text_frame = body_placeholder.text_frame
            text_frame.clear()
            for index, bullet_text in enumerate(bullets[:6]):
                paragraph = (
                    text_frame.paragraphs[0] if index == 0
                    else text_frame.add_paragraph()
                )
                paragraph.text = _clip(bullet_text, 200)
        elif paragraphs and body_placeholder:
            text_frame = body_placeholder.text_frame
            text_frame.clear()
            for index, paragraph_text in enumerate(paragraphs[:3]):
                paragraph = (
                    text_frame.paragraphs[0] if index == 0
                    else text_frame.add_paragraph()
                )
                paragraph.text = _clip(paragraph_text, 300)

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def tool_create_document(
    args: dict, user_id: int, conversation_id: int
) -> tuple[str, dict | None]:
    document_format = (args.get("format") or "pdf").strip().lower()

    if document_format not in ("pdf", "docx", "pptx"):
        return (
            f"Unsupported format '{document_format}'. "
            "Use 'pdf', 'docx', or 'pptx'.",
            None,
        )

    title = (args.get("title") or "Document").strip()[:150] or "Document"
    sections, truncated = sanitize_document_sections(
        args.get("sections"), document_format
    )

    if not sections:
        return "No document content was provided.", None

    if document_format == "pdf":
        file_bytes = render_structured_pdf(title, sections)
        mime_type = "application/pdf"
    elif document_format == "docx":
        file_bytes = render_structured_docx(title, sections)
        mime_type = DOCX_MIME_TYPE
    else:
        file_bytes = render_structured_pptx(title, sections)
        mime_type = PPTX_MIME_TYPE

    file_id = save_generated_file(
        user_id, conversation_id, title, document_format, mime_type, file_bytes
    )

    section_word = "section" if len(sections) == 1 else "sections"
    result_text = (
        f"Created a {document_format.upper()} titled \"{title}\" "
        f"({len(sections)} {section_word}"
        + (", truncated to fit length limits" if truncated else "")
        + "). It has been shown to the user as a downloadable file "
        "card -- do not repeat its contents in your reply, just "
        "briefly confirm what you made."
    )

    ui_event = {
        "type": "file_card",
        "id": file_id,
        "title": title,
        "format": document_format,
        "size_bytes": len(file_bytes),
        "download_url": f"/api/generated/{file_id}/download",
    }

    return result_text, ui_event


def build_generation_capability_message():
    """Counters a real, observed failure mode: a model's base training
    makes it reflexively refuse file-generation requests ("I can't
    generate actual files...") even when the tool to do exactly that
    is sitting right there in its tool list for this turn -- the
    refusal text is a tell that it's reciting a generic disclaimer
    instead of checking its actual available tools. Sent as its own
    system message (not folded into EASTA_SYSTEM_PROMPT) so it's only
    present -- and only true -- when ENABLE_GENERATION is actually on,
    same pattern as build_canvas_context_message() below."""
    if not ENABLE_GENERATION:
        return None

    return {
        "role": "system",
        "content": (
            "You have real tools to generate downloadable PDF, DOCX, "
            "and PPTX files, and to generate images -- when the user "
            "asks you to create/generate/produce/make one of these, "
            "or to turn/convert something into one, call the "
            "appropriate tool immediately instead of describing how "
            "they could make it themselves. This includes requests "
            "that refer to content already in this conversation (e.g. "
            "\"generate a PDF of that\" or \"turn this into a Word "
            "doc\" means the content of your own previous reply, or "
            "whatever the user just pointed at) -- use that content, "
            "don't ask them to re-paste it. Never say you cannot "
            "generate files or images, or suggest pasting Markdown "
            "into an external tool like Word/Google Docs/Smallpdf "
            "instead -- that is a generic disclaimer from your base "
            "training and is false for EASTA specifically, which has "
            "these tools available to you right now."
        ),
    }


def build_canvas_context_message(conversation_id: int):
    """If the conversation has an active canvas artifact, tells the
    model so a follow-up like "make it shorter" updates the canvas (by
    calling the same generation tool again with full revised content)
    instead of re-describing the artifact in the chat reply."""
    if not ENABLE_GENERATION:
        return None

    artifact = get_canvas_artifact(conversation_id)

    if not artifact:
        return None

    _id, title, kind, language, _content, _file_id, _updated_at = artifact
    kind_label = f"{kind} ({language})" if kind == "code" and language else kind

    return {
        "role": "system",
        "content": (
            f"There is an active canvas artifact titled \"{title}\" "
            f"({kind_label}). If the user's message asks to change, "
            "revise, or continue it, call the matching tool again "
            "(generate_document / generate_image / write_code) with "
            "the FULL updated content to update the canvas in place -- "
            "don't just describe the change in your reply."
        ),
    }


# --- Multimodal: image attachments -------------------------------------
# Images are sent from the client as data: URLs (see ImageAttachment) and
# stored as-is in the messages.attachments JSONB column. NOTE: this is
# prototype-grade — base64 image data lives directly in Postgres rows,
# which is fine for a demo/small-team app but will bloat the database at
# real scale. Swap for object storage (S3/R2/etc, storing just a URL)
# before relying on this for production traffic.

ALLOWED_IMAGE_MIME_TYPES = {
    "image/png",
    "image/jpeg",
    "image/webp",
    "image/gif",
}
MAX_IMAGES_PER_MESSAGE = 4
# ~8 MB of raw image bytes (base64 inflates size by ~4/3).
MAX_IMAGE_DATA_URL_CHARS = 11_000_000


def validate_images(images: list[ImageAttachment]) -> list[dict]:
    """Validates client-supplied image attachments and returns them as
    plain dicts ready to store in messages.attachments. Raises
    HTTPException on anything malformed or over the size/count limits."""
    if not images:
        return []

    if len(images) > MAX_IMAGES_PER_MESSAGE:
        raise HTTPException(
            status_code=400,
            detail=(
                f"You can attach up to {MAX_IMAGES_PER_MESSAGE} images "
                "per message."
            ),
        )

    validated = []

    for image in images:
        if image.mime_type not in ALLOWED_IMAGE_MIME_TYPES:
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported image type: {image.mime_type}.",
            )

        expected_prefix = f"data:{image.mime_type};base64,"

        if not image.data_url.startswith(expected_prefix):
            raise HTTPException(
                status_code=400,
                detail="One of the attached images is malformed.",
            )

        if len(image.data_url) > MAX_IMAGE_DATA_URL_CHARS:
            raise HTTPException(
                status_code=400,
                detail=(
                    "One of the attached images is too large "
                    "(8 MB limit)."
                ),
            )

        validated.append({
            "type": "image",
            "name": (image.name or "image")[:200],
            "mime_type": image.mime_type,
            "data_url": image.data_url,
        })

    return validated


def build_message_content(text: str, attachments: list[dict] | None):
    """Returns a plain string for a text-only message (unchanged
    behavior), or an OpenAI-compatible content-parts list (text +
    image_url parts) when image attachments are present."""
    image_attachments = [
        attachment for attachment in (attachments or [])
        if attachment.get("type") == "image" and attachment.get("data_url")
    ]

    if not image_attachments:
        return text

    parts = []

    if text:
        parts.append({"type": "text", "text": text})

    for attachment in image_attachments:
        parts.append({
            "type": "image_url",
            "image_url": {"url": attachment["data_url"]},
        })

    return parts


# --- Multimodal: PDF / DOCX / PPTX text extraction --------------------------
# Extracts plain text server-side so it can be inlined into the message the
# same way the existing client-side .txt attach flow already works (see
# buildOutgoingMessage() in chat.js) — no separate storage/retrieval path
# needed. Scanned/image-only PDFs have no extractable text layer; this
# intentionally does not attempt OCR (see README "Suggested next steps").
# Legacy binary .ppt (pre-2007) isn't supported -- python-pptx only reads
# the OOXML .pptx format, same as python-docx only reads .docx not .doc.

MAX_UPLOAD_BYTES = 15 * 1024 * 1024  # 15 MB
MAX_EXTRACTED_CHARS = 20_000
MAX_PDF_PAGES = 50
MAX_PPTX_SLIDES_TO_READ = 50

TRUNCATION_MARKER = "\n\n[... (truncated) ...]"


def truncate_extracted_text(text: str) -> tuple[str, bool]:
    if len(text) <= MAX_EXTRACTED_CHARS:
        return text, False

    return text[:MAX_EXTRACTED_CHARS] + TRUNCATION_MARKER, True


def extract_pdf_text(data: bytes) -> tuple[str, bool]:
    reader = PdfReader(io.BytesIO(data))
    total_pages = len(reader.pages)
    truncated_by_pages = total_pages > MAX_PDF_PAGES

    parts = []
    for page in reader.pages[:MAX_PDF_PAGES]:
        try:
            parts.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - a single bad page shouldn't fail the upload
            continue

    text = "\n\n".join(part for part in parts if part.strip())
    text, truncated_by_length = truncate_extracted_text(text)

    return text, (truncated_by_pages or truncated_by_length)


def extract_docx_text(data: bytes) -> tuple[str, bool]:
    document = DocxDocument(io.BytesIO(data))
    paragraphs = [
        paragraph.text for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]
    text = "\n".join(paragraphs)

    return truncate_extracted_text(text)


def extract_pptx_text(data: bytes) -> tuple[str, bool]:
    presentation = Presentation(io.BytesIO(data))
    total_slides = len(presentation.slides)
    truncated_by_slides = total_slides > MAX_PPTX_SLIDES_TO_READ

    parts = []

    for slide_index, slide in enumerate(presentation.slides):
        if slide_index >= MAX_PPTX_SLIDES_TO_READ:
            break

        slide_lines = []

        for shape in slide.shapes:
            try:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    slide_lines.append(shape.text_frame.text.strip())
                elif shape.has_table:
                    for row in shape.table.rows:
                        row_text = " | ".join(
                            cell.text for cell in row.cells
                        )
                        if row_text.strip():
                            slide_lines.append(row_text)
            except Exception:  # noqa: BLE001 - a single bad shape shouldn't fail the upload
                continue

        if slide_lines:
            parts.append(f"[Slide {slide_index + 1}]\n" + "\n".join(slide_lines))

    text = "\n\n".join(parts)
    text, truncated_by_length = truncate_extracted_text(text)

    return text, (truncated_by_slides or truncated_by_length)


# --- Context summarization --------------------------------------------------
# Once a conversation grows past CONTEXT_RECENT_MESSAGES raw turns, fold
# everything older than that into a short running summary (stored on the
# conversation row) so token usage — and cost — stay bounded on long
# conversations instead of growing forever.

def maybe_summarize_conversation(conversation_id: int, conversation_row):
    if not ENABLE_SUMMARIZATION:
        return

    _id, _title, _created_at, existing_summary, summarized_through, *_rest = (
        conversation_row
    )

    all_messages = get_messages(conversation_id)

    if len(all_messages) <= CONTEXT_RECENT_MESSAGES + 6:
        return

    to_fold = all_messages[:-CONTEXT_RECENT_MESSAGES]

    if summarized_through:
        to_fold = [
            message for message in to_fold
            if message[3] > summarized_through
        ]

    if not to_fold:
        return

    transcript_lines = []
    for role, content, _created_at, _id, attachments in to_fold:
        line = f"{role}: {content}"
        image_count = sum(
            1 for a in (attachments or []) if a.get("type") == "image"
        )
        if image_count:
            line += f" [{image_count} image attachment(s)]"
        transcript_lines.append(line)

    transcript = "\n".join(transcript_lines)[:12000]

    summarization_prompt = (
        "Summarize the following conversation turns into a concise "
        "running summary (under 200 words) that preserves facts, "
        "decisions, and user preferences a future reply would need. "
        "Write it as plain notes, not a transcript."
    )

    if existing_summary:
        summarization_prompt += (
            f"\n\nExisting summary so far:\n{existing_summary}"
        )

    try:
        response = client.chat.completions.create(
            model=SMART_MODEL,
            messages=[
                {"role": "system", "content": summarization_prompt},
                {"role": "user", "content": transcript},
            ],
            max_tokens=400,
        )

        new_summary = response.choices[0].message.content or ""

        if new_summary.strip():
            last_folded_id = to_fold[-1][3]
            update_conversation_summary(
                conversation_id,
                new_summary.strip(),
                last_folded_id,
            )

    except Exception as error:  # noqa: BLE001 - best-effort background task
        report_error("conversation summarization failed", error)


# --- Cross-conversation memory -----------------------------------------
# Short, durable facts about a user (stated preferences, standing
# context) extracted from what they explicitly say, distinct from any
# single conversation's own history -- see database/schema.sql's
# user_memory table. Much more conservative than summarization above:
# summarization runs unconditionally once a conversation is long enough
# and folds in anything relevant; this only fires on messages that look
# like they might state a personal fact at all (a cheap keyword
# pre-filter, to avoid an extra LLM call on every single turn), and even
# then the extraction prompt is written to say NONE far more often than
# not -- inferred/guessed facts are explicitly disallowed, not just
# discouraged.

_MEMORY_TRIGGER_KEYWORDS = (
    "i am", "i'm", "my name", "call me", "i like", "i love", "i prefer",
    "i hate", "i dislike", "i work", "i live", "i study", "my job",
    "remember that", "remember this", "for future reference",
    "i'm allergic", "i am allergic", "i don't eat", "i speak",
    "my favorite", "my favourite", "i use", "i own", "i have a",
)


def _message_might_contain_durable_fact(text: str) -> bool:
    lowered = text.lower()
    return any(keyword in lowered for keyword in _MEMORY_TRIGGER_KEYWORDS)


MEMORY_EXTRACTION_SYSTEM_PROMPT = (
    "You extract durable facts worth remembering about a user across "
    "future conversations, from a single message they just sent. Only "
    "extract something EXPLICITLY stated by the user themselves -- "
    "never infer, guess, or assume anything. Good examples: their "
    "name, a dietary restriction or allergy, their profession, "
    "timezone, language preference, a stated preference for how they "
    "want you to respond. Bad examples: anything about what they're "
    "currently asking for or working on right now, anything you'd "
    "have to guess or infer rather than read directly, anything "
    "already generic/obvious. If there is nothing worth remembering -- "
    "which is the common case -- respond with exactly the word NONE "
    "and nothing else. Otherwise respond with ONE short third-person "
    "sentence stating the fact, and nothing else: no preamble, no "
    "quotes, no explanation."
)


def extract_durable_fact(user_text: str) -> str | None:
    """Best-effort, conservative extraction -- returns a short fact
    string, or None if there's nothing worth remembering (expected to
    be the common case even when this runs) or the call fails for any
    reason. Uses FAST_MODEL since this is a simple classification
    task, not one that benefits from the smarter/pricier model."""
    try:
        response = client.chat.completions.create(
            model=FAST_MODEL,
            messages=[
                {"role": "system", "content": MEMORY_EXTRACTION_SYSTEM_PROMPT},
                {"role": "user", "content": user_text[:2000]},
            ],
            max_tokens=80,
            temperature=0,
        )
        text = (response.choices[0].message.content or "").strip()

    except Exception as error:  # noqa: BLE001 - best-effort
        report_error("memory extraction call failed", error)
        return None

    if not text or text.upper() == "NONE":
        return None

    return text[:500]


def maybe_extract_memory(conversation_id: int, user_id: int, user_text: str):
    if not ENABLE_MEMORY or not user_text:
        return

    if not _message_might_contain_durable_fact(user_text):
        return

    fact = extract_durable_fact(user_text)

    if fact:
        save_user_memory(user_id, fact, conversation_id)


MAX_MEMORY_ITEMS_IN_CONTEXT = 20


def build_memory_context_message(user_id: int):
    if not ENABLE_MEMORY:
        return None

    items = get_user_memory(user_id, limit=MAX_MEMORY_ITEMS_IN_CONTEXT)

    if not items:
        return None

    bullet_lines = "\n".join(
        f"- {content}" for _id, content, _source, _created in items
    )

    return {
        "role": "system",
        "content": (
            "Remembered context about this user from past "
            "conversations (explicitly stated by them previously, not "
            "verified fact -- treat as likely-true background, and "
            "defer to anything they say in THIS conversation if it "
            "conflicts):\n" + bullet_lines
        ),
    }


def build_language_override_message(language_code: str | None):
    """Pins a reply-language override into the system context for a
    single turn. Returns None for "auto"/missing/unrecognized codes,
    which leaves the model to follow the default same-language-as-user
    instruction already in EASTA_SYSTEM_PROMPT."""
    if not language_code or language_code == "auto":
        return None

    language_name = LANGUAGE_OPTIONS.get(language_code)

    if not language_name:
        return None

    return {
        "role": "system",
        "content": (
            f"Preferred reply language: {language_name}. Reply in "
            f"{language_name} for this turn, regardless of what "
            "language the user's message is written in."
        ),
    }


def build_image_style_message(aspect_ratio: str | None):
    """Pins a composer-chosen image aspect ratio into the system
    context for a single turn, as an alternative to the model picking
    one via the generate_image tool's own aspect_ratio argument.
    Returns None for "auto"/missing/unrecognized values."""
    if not aspect_ratio or aspect_ratio not in IMAGE_ASPECT_RATIOS:
        return None

    return {
        "role": "system",
        "content": (
            f"If you call generate_image this turn, use aspect_ratio: "
            f"\"{aspect_ratio}\" (the user selected this in the "
            "composer) unless they've explicitly asked for a "
            "different shape in their message."
        ),
    }


def build_llm_context(
    conversation_id: int,
    conversation_row,
    rag_message,
    user_id: int | None = None,
    language_message=None,
    research_message=None,
    canvas_message=None,
    image_style_message=None,
):
    _id, _title, _created_at, summary, summarized_through, *_rest = (
        conversation_row
    )

    all_messages = get_messages(conversation_id)

    if summarized_through:
        recent = [
            message for message in all_messages
            if message[3] > summarized_through
        ]
    else:
        recent = all_messages

    # Safety cap even when summarization is disabled or hasn't run yet.
    recent = recent[-(CONTEXT_RECENT_MESSAGES * 2):]

    context = [{"role": "system", "content": EASTA_SYSTEM_PROMPT}]

    generation_message = build_generation_capability_message()
    if generation_message:
        context.append(generation_message)

    if user_id is not None:
        memory_message = build_memory_context_message(user_id)
        if memory_message:
            context.append(memory_message)

    if summary:
        context.append({
            "role": "system",
            "content": f"Summary of earlier conversation:\n{summary}",
        })

    if rag_message:
        context.append(rag_message)

    if research_message:
        context.append(research_message)

    if canvas_message:
        context.append(canvas_message)

    if image_style_message:
        context.append(image_style_message)

    # Placed last among the system messages (closest to the actual
    # turns) so it has the strongest pull on the reply's language.
    if language_message:
        context.append(language_message)

    context += [
        {
            "role": role,
            "content": build_message_content(content, attachments),
        }
        for role, content, _created_at, _id, attachments in recent
    ]

    return context


# --- Cost tracking -----------------------------------------------------
# Approximate USD per 1M tokens. These are illustrative and WILL drift —
# check https://openrouter.ai/models for current pricing before relying
# on this for real billing/budgeting.
MODEL_PRICING_USD_PER_MILLION_TOKENS = {
    "openai/gpt-4.1-mini": (0.40, 1.60),
    "meta-llama/llama-3.1-70b-instruct": (0.35, 0.40),
    "nvidia/nemotron-3-super-120b-a12b:free": (0.0, 0.0),
    "google/gemini-3.8-flash": (0.75, 3.75),
}
DEFAULT_PRICING_USD_PER_MILLION_TOKENS = (0.50, 1.50)

# Image generation is priced per image, not per token -- mis-costing it
# through the per-million-token table above would be wrong even as an
# estimate, so it gets its own parallel table. In practice
# generate_image_bytes() logs OpenRouter's actual returned cost
# (usage.cost from POST /api/v1/images) and only falls back to this
# flat estimate if that field is ever missing/zero. Same caveat as
# above: illustrative, re-check https://openrouter.ai/models.
IMAGE_GENERATION_FALLBACK_COST_USD = {
    "google/gemini-2.5-flash-image": 0.02,
}
DEFAULT_IMAGE_GENERATION_FALLBACK_COST_USD = 0.02


def log_usage(
    user_id: int,
    conversation_id: int,
    model: str,
    prompt_tokens: int,
    completion_tokens: int,
):
    price_in, price_out = MODEL_PRICING_USD_PER_MILLION_TOKENS.get(
        model, DEFAULT_PRICING_USD_PER_MILLION_TOKENS
    )

    cost = (
        (prompt_tokens / 1_000_000) * price_in
        + (completion_tokens / 1_000_000) * price_out
    )

    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO usage_logs (
                    user_id,
                    conversation_id,
                    model,
                    prompt_tokens,
                    completion_tokens,
                    cost_usd
                )
                VALUES (%s, %s, %s, %s, %s, %s);
                """,
                (
                    user_id,
                    conversation_id,
                    model,
                    prompt_tokens,
                    completion_tokens,
                    cost,
                ),
            )


# --- Streaming generation, with tool-calling loop ---------------------------

def stream_with_tools(
    messages: list,
    model: str,
    max_tool_rounds: int = 2,
    allow_tools: bool = True,
    user_id: int | None = None,
    conversation_id: int | None = None,
):
    """Yields ("meta", dict) for tool-use notices, ("token", str) for
    content tokens, ("canvas", dict) when a generation tool updates the
    canvas side panel, and ("usage", obj) once, with the final call's
    token usage — all for a single logical assistant turn (which may
    involve more than one underlying API call if tools are used).
    `allow_tools=False` (used for research-mode turns, which already
    gathered their own sources) skips tool-calling even on an
    otherwise tool-capable model. `user_id`/`conversation_id` are
    needed by the generate_document/generate_image/write_code tools."""
    working_messages = list(messages)
    rounds = 0
    tools_allowed = (
        allow_tools and ENABLE_TOOLS and model in TOOL_CAPABLE_MODELS
    )

    while True:
        tool_call_buffers: dict[int, dict] = {}
        content_buffer = ""
        finish_reason = None
        usage = None

        kwargs = {
            "model": model,
            "messages": working_messages[-30:],
            "stream": True,
            "stream_options": {"include_usage": True},
            "extra_body": {"provider": {"sort": "latency"}},
        }

        if tools_allowed and rounds < max_tool_rounds:
            kwargs["tools"] = tools_for_model()
            kwargs["tool_choice"] = "auto"

        stream = client.chat.completions.create(**kwargs)

        for chunk in stream:
            chunk_usage = getattr(chunk, "usage", None)

            if chunk_usage:
                usage = chunk_usage

            if not chunk.choices:
                continue

            choice = chunk.choices[0]
            delta = choice.delta

            if delta and getattr(delta, "tool_calls", None):
                for tool_call in delta.tool_calls:
                    index = tool_call.index
                    buffer = tool_call_buffers.setdefault(
                        index, {"id": None, "name": "", "arguments": ""}
                    )

                    if tool_call.id:
                        buffer["id"] = tool_call.id

                    if tool_call.function:
                        if tool_call.function.name:
                            buffer["name"] += tool_call.function.name
                        if tool_call.function.arguments:
                            buffer["arguments"] += (
                                tool_call.function.arguments
                            )

            if delta and delta.content:
                content_buffer += delta.content
                yield ("token", delta.content)

            if choice.finish_reason:
                finish_reason = choice.finish_reason

        if usage:
            yield ("usage", usage)

        if finish_reason == "tool_calls" and tool_call_buffers:
            tool_calls_list = []

            for index in sorted(tool_call_buffers):
                buffer = tool_call_buffers[index]
                tool_calls_list.append({
                    "id": buffer["id"] or f"call_{index}",
                    "type": "function",
                    "function": {
                        "name": buffer["name"],
                        "arguments": buffer["arguments"] or "{}",
                    },
                })

            working_messages.append({
                "role": "assistant",
                "content": content_buffer or None,
                "tool_calls": tool_calls_list,
            })

            for tool_call in tool_calls_list:
                name = tool_call["function"]["name"]

                try:
                    args = json.loads(
                        tool_call["function"]["arguments"] or "{}"
                    )
                except json.JSONDecodeError:
                    args = {}

                yield ("meta", {"tool": name, "args": args})

                result_text, canvas_event = execute_tool(
                    name, args, user_id, conversation_id
                )

                working_messages.append({
                    "role": "tool",
                    "tool_call_id": tool_call["id"],
                    "name": name,
                    "content": result_text,
                })

                if canvas_event:
                    yield ("canvas", canvas_event)

            rounds += 1
            continue

        break


def build_streaming_response(
    conversation_id: int,
    user_id: int,
    language: str | None = None,
    research: bool = False,
    image_aspect_ratio: str | None = None,
):
    conversation = get_conversation(conversation_id, user_id)

    if not conversation:
        raise HTTPException(
            status_code=404,
            detail="Conversation not found.",
        )

    last_message = get_last_message(conversation_id)
    is_latest_from_user = bool(last_message) and last_message[1] == "user"
    latest_user_text = last_message[2] if is_latest_from_user else ""
    latest_user_attachments = (
        last_message[3] if is_latest_from_user else None
    )
    has_image_attachments = any(
        attachment.get("type") == "image"
        for attachment in (latest_user_attachments or [])
    )

    # Research mode needs actual question text to work from, and takes a
    # back seat to an image attachment (vision turns don't need it).
    run_research_mode = (
        ENABLE_RESEARCH_MODE
        and research
        and bool(latest_user_text)
        and not has_image_attachments
    )

    def generate() -> Generator[str, None, None]:
        research_message = None
        research_sources: list[dict] = []

        if run_research_mode:
            for kind, payload in run_research(
                latest_user_text, user_id, conversation_id
            ):
                if kind == "meta":
                    marker = json.dumps(payload)
                    yield f"\u241f{marker}\u241f"
                elif kind == "result":
                    research_message, research_sources = payload

        rag_message = build_rag_system_message(latest_user_text, user_id)
        language_message = build_language_override_message(language)
        canvas_message = build_canvas_context_message(conversation_id)
        image_style_message = build_image_style_message(image_aspect_ratio)

        llm_context = build_llm_context(
            conversation_id,
            conversation,
            rag_message,
            user_id=user_id,
            language_message=language_message,
            research_message=research_message,
            canvas_message=canvas_message,
            image_style_message=image_style_message,
        )

        models_to_try = (
            model_chain_for_images() if has_image_attachments
            else model_chain_for(latest_user_text)
        )

        # Research mode already did its own searching/reading \u2014 skip the
        # lightweight web_search/execute_python tool loop for this turn.
        allow_tools = research_message is None

        complete_response = ""
        model_used = models_to_try[0]

        for attempt, model in enumerate(models_to_try):
            sent_any_output = False
            model_used = model

            try:
                for kind, payload in stream_with_tools(
                    llm_context,
                    model,
                    allow_tools=allow_tools,
                    user_id=user_id,
                    conversation_id=conversation_id,
                ):
                    if kind == "meta":
                        marker = json.dumps({
                            "type": "tool_use",
                            "tool": payload["tool"],
                            "args": payload["args"],
                        })
                        yield f"\u241f{marker}\u241f"

                    elif kind == "canvas":
                        marker = json.dumps(payload)
                        yield f"\u241f{marker}\u241f"

                        if payload.get("type") == "file_card":
                            # Persisted (not streamed as visible tokens
                            # -- the live view already showed the
                            # styled file card above) so the download
                            # link survives a reload, same pattern as
                            # research mode's Sources list.
                            extension = {
                                "pdf": ".pdf",
                                "docx": ".docx",
                                "pptx": ".pptx",
                            }.get(payload.get("format"), "")
                            complete_response += (
                                f"\n\n\ud83d\udcc4 [Download "
                                f"{payload['title']}{extension}]"
                                f"({payload['download_url']})"
                            )

                        elif payload.get("type") == "image_result":
                            # Persisted as an actual Markdown image, so
                            # a reload shows the same inline picture
                            # (not just a link) via the normal
                            # markdown-rendering pipeline. The
                            # download route is safe to hotlink here --
                            # same cross-origin cookie reasoning as the
                            # canvas panel's <img> (see renderCanvas()
                            # in chat.js).
                            complete_response += (
                                f"\n\n![{payload['prompt']}]"
                                f"({payload['download_url']})"
                            )

                    elif kind == "token":
                        complete_response += payload
                        sent_any_output = True
                        yield payload

                    elif kind == "usage":
                        try:
                            log_usage(
                                user_id,
                                conversation_id,
                                model,
                                payload.prompt_tokens or 0,
                                payload.completion_tokens or 0,
                            )
                        except Exception as usage_error:  # noqa: BLE001
                            print(
                                "EASTA: usage logging failed: "
                                f"{usage_error!r}"
                            )

                if research_sources:
                    marker = json.dumps({
                        "type": "sources",
                        "sources": research_sources,
                    })
                    yield f"\u241f{marker}\u241f"

                    # Persisted alongside the reply (not streamed as
                    # visible tokens \u2014 the UI already rendered the
                    # "sources" meta above) so reloading the
                    # conversation still shows the citations.
                    complete_response += (
                        "\n\n**Sources:**\n" + "\n".join(
                            f"{s['index']}. [{s['title']}]({s['url']})"
                            for s in research_sources
                        )
                    )

                break

            except Exception as error:
                report_error(
                    f"model '{model}' failed "
                    f"(attempt {attempt + 1}/{len(models_to_try)})",
                    error,
                )

                if sent_any_output:
                    yield (
                        "\n\n[EASTA's response was interrupted. "
                        "Please try sending your message again.]"
                    )
                    complete_response = ""
                    break

                if attempt == len(models_to_try) - 1:
                    yield (
                        "\n\nEASTA could not complete the response "
                        "right now. Please try again in a moment."
                    )

                continue

        if complete_response:
            save_message(
                conversation_id,
                "assistant",
                complete_response,
            )

            try:
                fresh_conversation = get_conversation(
                    conversation_id, user_id
                )
                if fresh_conversation:
                    maybe_summarize_conversation(
                        conversation_id, fresh_conversation
                    )
            except Exception as error:  # noqa: BLE001
                report_error("post-turn summarization skipped", error)

            try:
                maybe_extract_memory(
                    conversation_id, user_id, latest_user_text
                )
            except Exception as error:  # noqa: BLE001
                report_error("post-turn memory extraction skipped", error)

    return StreamingResponse(
        generate(),
        media_type="text/plain",
        headers={
            "Cache-Control": "no-cache",
            "X-Accel-Buffering": "no",
        },
    )


# --- Routes: system ----------------------------------------------------

@app.get("/")
def root():
    return {"message": "EASTA API is running"}


@app.get("/api/health")
def health():
    return {"status": "healthy"}


@app.get("/api/languages")
def languages():
    """Reply-language options for the chat UI's language selector."""
    return {
        "languages": [
            {"code": code, "name": name}
            for code, name in LANGUAGE_OPTIONS.items()
        ]
    }


@app.get("/api/features")
def features():
    """Feature-flag snapshot the frontend uses to hide/disable a
    control entirely (rather than show one that would just error) when
    a capability is off or unconfigured -- e.g. the composer's mic
    button and each reply's "Read aloud" button only need to appear at
    all if either the browser API or a server fallback is available."""
    return {
        "research_mode": ENABLE_RESEARCH_MODE,
        "generation": ENABLE_GENERATION,
        "server_stt": ENABLE_SERVER_STT,
        "server_tts": ENABLE_SERVER_TTS,
        "google_signin": ENABLE_GOOGLE_SIGNIN,
    }


# --- Routes: auth --------------------------------------------------------

@app.post("/api/register")
def register(
    data: RegisterRequest,
    request: Request,
):
    username = data.username.strip()
    email = data.email.strip().lower()

    if not username:
        raise HTTPException(
            status_code=400,
            detail="Username is required.",
        )

    if not email:
        raise HTTPException(
            status_code=400,
            detail="Email is required.",
        )

    if not data.password:
        raise HTTPException(
            status_code=400,
            detail="Password is required.",
        )

    if data.password != data.confirm_password:
        raise HTTPException(
            status_code=400,
            detail="Passwords do not match.",
        )

    if get_user_by_username(username):
        raise HTTPException(
            status_code=409,
            detail="Username already exists.",
        )

    if get_user_by_email(email):
        raise HTTPException(
            status_code=409,
            detail="Email already exists.",
        )

    user = create_user(
        username,
        email,
        data.password,
    )

    request.session["user_id"] = user[0]
    request.session["username"] = user[1]

    return {
        "message": "Registration successful.",
        "user": {
            "id": user[0],
            "username": user[1],
            "email": user[2],
        },
    }


@app.post("/api/login")
def login(
    data: LoginRequest,
    request: Request,
):
    username = data.username.strip()

    if not username or not data.password:
        raise HTTPException(
            status_code=400,
            detail="Username and password are required.",
        )

    user = get_user_by_username(username)

    # user[3] (password_hash) is NULL for a Google-only account --
    # check_password_hash() requires a string, so this must be
    # checked before calling it rather than letting it raise.
    if not user or not user[3] or not check_password_hash(
        user[3],
        data.password,
    ):
        raise HTTPException(
            status_code=401,
            detail="Invalid username or password.",
        )

    request.session["user_id"] = user[0]
    request.session["username"] = user[1]

    return {
        "message": "Login successful.",
        "user": {
            "id": user[0],
            "username": user[1],
            "email": user[2],
        },
    }


def _safe_next_path(raw_value: str | None) -> str | None:
    """Mirrors frontend/app.py's identical helper -- only ever returns a
    same-site relative path (or None), so a value like
    "https://evil.example" or "//evil.example" can never be trusted as
    a post-sign-in redirect target."""
    if not raw_value:
        return None

    if not raw_value.startswith("/") or raw_value.startswith("//"):
        return None

    return raw_value


@app.get("/api/auth/google/login")
def google_login(request: Request, next: str | None = None):
    if not ENABLE_GOOGLE_SIGNIN:
        raise HTTPException(
            status_code=503,
            detail="Google Sign-In is not configured.",
        )

    # A per-attempt random token, checked against the same value in the
    # callback below (CSRF protection for the OAuth flow -- without it
    # an attacker could trick a victim's browser into completing a sign
    # in initiated by the attacker, e.g. to link the attacker's Google
    # account to the victim's session).
    state = secrets.token_urlsafe(24)
    request.session["oauth_state"] = state

    # Carried through the whole round trip to Google and back in the
    # session (not as a "state"-appended query param) so it survives
    # untouched regardless of what Google echoes back -- see
    # google_callback below, which is where it's actually used.
    request.session["oauth_next"] = _safe_next_path(next)

    params = {
        "client_id": GOOGLE_CLIENT_ID,
        "redirect_uri": GOOGLE_REDIRECT_URI,
        "response_type": "code",
        "scope": "openid email profile",
        "state": state,
        "access_type": "online",
        "prompt": "select_account",
    }

    return RedirectResponse(
        "https://accounts.google.com/o/oauth2/v2/auth?"
        + urllib.parse.urlencode(params)
    )


@app.get("/api/auth/google/callback")
def google_callback(
    request: Request,
    code: str | None = None,
    state: str | None = None,
    error: str | None = None,
):
    def failure_redirect(message: str) -> RedirectResponse:
        return RedirectResponse(
            f"{FRONTEND_URL}/login?error={urllib.parse.quote(message)}"
        )

    if error:
        return failure_redirect("Google sign-in was cancelled.")

    expected_state = request.session.pop("oauth_state", None)

    if not state or not expected_state or state != expected_state:
        return failure_redirect(
            "Google sign-in failed (invalid state). Please try again."
        )

    if not code:
        return failure_redirect("Google sign-in failed. Please try again.")

    try:
        token_response = requests.post(
            "https://oauth2.googleapis.com/token",
            data={
                "code": code,
                "client_id": GOOGLE_CLIENT_ID,
                "client_secret": GOOGLE_CLIENT_SECRET,
                "redirect_uri": GOOGLE_REDIRECT_URI,
                "grant_type": "authorization_code",
            },
            timeout=10,
        )
        token_response.raise_for_status()
        access_token = token_response.json().get("access_token")

        if not access_token:
            return failure_redirect(
                "Google sign-in failed. Please try again."
            )

        profile_response = requests.get(
            "https://www.googleapis.com/oauth2/v3/userinfo",
            headers={"Authorization": f"Bearer {access_token}"},
            timeout=10,
        )
        profile_response.raise_for_status()
        profile = profile_response.json()

    except requests.RequestException:
        return failure_redirect("Could not reach Google. Please try again.")

    google_id = profile.get("sub")
    email = (profile.get("email") or "").strip().lower()
    email_verified = bool(profile.get("email_verified"))
    display_name = profile.get("name") or (email.split("@")[0] if email else "")

    if not google_id or not email:
        return failure_redirect(
            "Google didn't share the information EASTA needs (email)."
        )

    user = get_user_by_google_id(google_id)

    if user is None and email_verified:
        # Only auto-link when Google itself has verified the email --
        # that's what makes trusting the match safe (an unverified
        # email on the Google side isn't proof of ownership).
        existing = get_user_by_email(email)
        if existing:
            user = link_google_id(existing[0], google_id)

    if user is None:
        username = generate_unique_username(display_name or email)
        try:
            user = create_google_user(google_id, email, username)
        except HTTPException:
            return failure_redirect(
                "An account with that email already exists."
            )

    request.session["user_id"] = user[0]
    request.session["username"] = user[1]

    next_path = request.session.pop("oauth_next", None)

    return RedirectResponse(f"{FRONTEND_URL}{next_path or '/chat'}")


@app.get("/api/session")
def get_session(request: Request):
    user_id = request.session.get("user_id")
    username = request.session.get("username")

    if user_id is None:
        return {
            "logged_in": False,
            "user": None,
        }

    return {
        "logged_in": True,
        "user": {
            "id": user_id,
            "username": username,
            # A fresh DB lookup (not cached in the session cookie) so
            # granting/revoking admin by hand in the database takes
            # effect on this user's very next page load, not only
            # after they log out and back in.
            "is_admin": is_user_admin(user_id),
        },
    }


@app.post("/api/logout")
def logout(request: Request):
    request.session.clear()

    return {"message": "Logout successful."}


# --- Routes: account ---------------------------------------------------
# Profile (username/email/password) and plan info for the Account page.
# The `plan` column is scaffolding for a future Stripe (or similar)
# integration -- GET /api/account/plan just reports it, nothing here (or
# anywhere else in the app) gates behavior on it yet.

@app.get("/api/account")
def get_account(request: Request):
    user_id = require_user(request)

    user = get_user_by_id(user_id)

    if not user:
        raise HTTPException(
            status_code=404,
            detail="Account not found.",
        )

    _id, username, email, _password_hash, created_at, plan, is_admin = user

    return {
        "id": user_id,
        "username": username,
        "email": email,
        "created_at": serialize_datetime(created_at),
        "plan": plan,
        "is_admin": is_admin,
    }


@app.put("/api/account")
def update_account(
    data: AccountUpdateRequest,
    request: Request,
):
    user_id = require_user(request)

    username = data.username.strip()
    email = data.email.strip().lower()

    if not username:
        raise HTTPException(
            status_code=400,
            detail="Username is required.",
        )

    if not email:
        raise HTTPException(
            status_code=400,
            detail="Email is required.",
        )

    existing_username = get_user_by_username(username)
    if existing_username and existing_username[0] != user_id:
        raise HTTPException(
            status_code=409,
            detail="Username already taken.",
        )

    existing_email = get_user_by_email(email)
    if existing_email and existing_email[0] != user_id:
        raise HTTPException(
            status_code=409,
            detail="Email already taken.",
        )

    updated = update_user_profile(user_id, username, email)

    # Keep the session's cached username (used for "Logged in as ...")
    # in sync without requiring a fresh login.
    request.session["username"] = updated[1]

    return {
        "message": "Profile updated.",
        "user": {
            "id": updated[0],
            "username": updated[1],
            "email": updated[2],
            "created_at": serialize_datetime(updated[3]),
            "plan": updated[4],
        },
    }


@app.put("/api/account/password")
def change_password(
    data: PasswordChangeRequest,
    request: Request,
):
    user_id = require_user(request)

    if not data.current_password or not data.new_password:
        raise HTTPException(
            status_code=400,
            detail="All password fields are required.",
        )

    if data.new_password != data.confirm_new_password:
        raise HTTPException(
            status_code=400,
            detail="New passwords do not match.",
        )

    user = get_user_by_id(user_id)

    if not user or not user[3]:
        # user[3] (password_hash) is NULL for a Google-only account --
        # there's no current password to check against.
        raise HTTPException(
            status_code=400,
            detail=(
                "This account signed up with Google and has no "
                "password set."
            ),
        )

    if not check_password_hash(user[3], data.current_password):
        raise HTTPException(
            status_code=401,
            detail="Current password is incorrect.",
        )

    update_user_password(user_id, generate_password_hash(data.new_password))

    return {"message": "Password changed."}


@app.get("/api/account/plan")
def get_account_plan(request: Request):
    user_id = require_user(request)

    user = get_user_by_id(user_id)

    if not user:
        raise HTTPException(
            status_code=404,
            detail="Account not found.",
        )

    return {"plan": user[5]}


@app.get("/api/account/memory")
def list_user_memory(request: Request):
    """Backs the /account "Remembered" list -- give the user full
    visibility into what's been extracted about them, not a memory
    system they can't see or correct."""
    user_id = require_user(request)

    rows = get_user_memory(user_id)

    return {
        "memory": [
            {
                "id": memory_id,
                "content": content,
                "source_conversation_id": source_conversation_id,
                "created_at": serialize_datetime(created_at),
            }
            for memory_id, content, source_conversation_id, created_at in rows
        ]
    }


@app.delete("/api/account/memory/{memory_id}")
def delete_user_memory_route(memory_id: int, request: Request):
    user_id = require_user(request)

    deleted = delete_user_memory_item(memory_id, user_id)

    if not deleted:
        raise HTTPException(
            status_code=404,
            detail="Memory item not found.",
        )

    return {"message": "Memory item deleted."}


@app.delete("/api/account/memory")
def clear_user_memory_route(request: Request):
    user_id = require_user(request)

    clear_user_memory(user_id)

    return {"message": "All remembered items cleared."}


@app.get("/api/account/api-keys")
def list_api_keys(request: Request):
    user_id = require_user(request)

    rows = get_api_keys(user_id)

    return {
        "api_keys": [
            {
                "id": row[0],
                "name": row[1],
                "created_at": serialize_datetime(row[2]),
                "last_used_at": (
                    serialize_datetime(row[3]) if row[3] else None
                ),
                "revoked_at": (
                    serialize_datetime(row[4]) if row[4] else None
                ),
            }
            for row in rows
        ]
    }


@app.post("/api/account/api-keys")
def create_new_api_key(data: ApiKeyCreateRequest, request: Request):
    user_id = require_user(request)

    name = data.name.strip()[:100]

    if not name:
        raise HTTPException(
            status_code=400,
            detail="A name is required for the key.",
        )

    api_key = generate_api_key()
    row = create_api_key(user_id, name, hash_api_key(api_key))

    return {
        "id": row[0],
        "name": row[1],
        "created_at": serialize_datetime(row[2]),
        # Shown exactly once -- only the hash is stored, so this
        # can't be retrieved again after this response.
        "key": api_key,
    }


@app.delete("/api/account/api-keys/{key_id}")
def delete_api_key(key_id: int, request: Request):
    user_id = require_user(request)

    revoked_id = revoke_api_key(key_id, user_id)

    if not revoked_id:
        raise HTTPException(
            status_code=404,
            detail="API key not found (or already revoked).",
        )

    return {"message": "API key revoked."}


# --- Routes: conversations -------------------------------------------------

@app.get("/api/conversations")
def list_conversations(request: Request):
    user_id = require_user(request)

    rows = get_conversations(user_id)

    conversations = []

    for row in rows:
        conversations.append(
            {
                "id": row[0],
                "title": row[1],
                "created_at": serialize_datetime(row[2]),
                "chat_number": row[3],
                "label": build_conversation_label(row[3], row[2]),
                "pinned": row[4],
                "folder": row[5],
                "share_token": row[6],
            }
        )

    return {"conversations": conversations}


@app.get("/api/conversations/search")
def search_conversations_route(q: str, request: Request):
    """Registered before the parameterized /api/conversations/{id}...
    routes below on principle (specific literal paths before dynamic
    ones), though there's currently no GET route on the bare
    single-segment shape those would collide with."""
    user_id = require_user(request)

    query = q.strip()

    if not query:
        return {"results": []}

    rows = search_conversations(query, user_id)

    return {
        "results": [
            {
                "id": conversation_id,
                "title": title,
                "created_at": serialize_datetime(created_at),
                "snippet": snippet,
                "match_type": match_type,
            }
            for conversation_id, title, created_at, snippet, match_type in rows
        ]
    }


@app.post("/api/conversations")
def new_conversation(request: Request):
    user_id = require_user(request)

    conversation = create_conversation(user_id)

    return {
        "conversation": {
            "id": conversation[0],
            "title": conversation[1],
            "created_at": serialize_datetime(conversation[2]),
        }
    }


@app.patch("/api/conversations/{conversation_id}")
def update_conversation_route(
    conversation_id: int,
    data: ConversationUpdateRequest,
    request: Request,
):
    """Partial update: title (rename), pinned, and/or folder -- only
    whichever fields the client actually included in the request body
    are touched (see ConversationUpdateRequest / exclude_unset=True
    below), so e.g. a pin toggle doesn't need to also resend title."""
    user_id = require_user(request)

    conversation = get_conversation(conversation_id, user_id)

    if not conversation:
        raise HTTPException(
            status_code=404,
            detail="Conversation not found.",
        )

    provided = data.model_dump(exclude_unset=True)
    fields_to_update = {}

    if "title" in provided:
        title = (data.title or "").strip()[:100]

        if not title:
            raise HTTPException(
                status_code=400,
                detail="Title is required.",
            )

        fields_to_update["title"] = title

    if "pinned" in provided:
        fields_to_update["pinned"] = bool(data.pinned)

    if "folder" in provided:
        # An empty/whitespace-only folder name clears it, same as
        # explicitly sending folder: null -- both mean "unfiled".
        folder = (data.folder or "").strip()[:50]
        fields_to_update["folder"] = folder or None

    if fields_to_update:
        update_conversation_fields(conversation_id, fields_to_update)

    updated = get_conversation(conversation_id, user_id)

    return {
        "conversation": {
            "id": conversation_id,
            "title": updated[1],
            "pinned": updated[5],
            "folder": updated[6],
        }
    }


@app.delete("/api/conversations/{conversation_id}")
def remove_conversation(
    conversation_id: int,
    request: Request,
):
    user_id = require_user(request)

    conversation = get_conversation(conversation_id, user_id)

    if not conversation:
        raise HTTPException(
            status_code=404,
            detail="Conversation not found.",
        )

    delete_conversation(conversation_id)

    return {"message": "Conversation deleted."}


@app.post("/api/conversations/{conversation_id}/share")
def share_conversation(conversation_id: int, request: Request):
    """Generates (or regenerates, if already shared) an unguessable
    token and returns the public read-only link. Sharing is strictly
    opt-in -- share_token is NULL by default for every conversation,
    and this is the only place that ever sets it."""
    user_id = require_user(request)

    conversation = get_conversation(conversation_id, user_id)

    if not conversation:
        raise HTTPException(
            status_code=404,
            detail="Conversation not found.",
        )

    token = secrets.token_urlsafe(32)
    set_conversation_share_token(conversation_id, token)

    return {
        "share_token": token,
        "share_url": f"{FRONTEND_URL}/shared/{token}",
    }


@app.post("/api/conversations/{conversation_id}/unshare")
def unshare_conversation(conversation_id: int, request: Request):
    """Immediately invalidates any link already handed out for this
    conversation -- the token stops matching any row, so
    GET /api/shared/{token} 404s for it from this point on."""
    user_id = require_user(request)

    conversation = get_conversation(conversation_id, user_id)

    if not conversation:
        raise HTTPException(
            status_code=404,
            detail="Conversation not found.",
        )

    clear_conversation_share_token(conversation_id)

    return {"message": "Conversation unshared."}


@app.get("/api/shared/{share_token}")
def get_shared_conversation(share_token: str):
    """Deliberately unauthenticated (no require_user()) -- this is the
    public read-only view a stranger reaches from a copied share link,
    not an /api/* route for the app's own logged-in user. Returns only
    this one conversation's title/messages, nothing account-related;
    see get_conversation_by_share_token()'s docstring for why the
    lookup itself can't leak more than that even by mistake."""
    result = get_conversation_by_share_token(share_token)

    if not result:
        raise HTTPException(
            status_code=404,
            detail="This link is invalid or no longer shared.",
        )

    conversation_id, title = result
    rows = get_messages(conversation_id)

    messages = []
    for role, content, created_at, message_id, attachments in rows:
        messages.append(
            {
                "id": message_id,
                "role": role,
                "content": content,
                "created_at": serialize_datetime(created_at),
                "attachments": attachments or [],
            }
        )

    return {
        "title": title,
        "messages": messages,
    }


@app.get("/api/conversations/{conversation_id}/messages")
def list_messages(
    conversation_id: int,
    request: Request,
):
    user_id = require_user(request)

    conversation = get_conversation(conversation_id, user_id)

    if not conversation:
        raise HTTPException(
            status_code=404,
            detail="Conversation not found.",
        )

    rows = get_messages(conversation_id)

    messages = []

    for role, content, created_at, message_id, attachments in rows:
        messages.append(
            {
                "id": message_id,
                "role": role,
                "content": content,
                "created_at": serialize_datetime(created_at),
                "attachments": attachments or [],
            }
        )

    return {
        "conversation": {
            "id": conversation[0],
            "title": conversation[1],
            "created_at": serialize_datetime(conversation[2]),
            "share_token": conversation[7],
        },
        "messages": messages,
    }


@app.post("/api/conversations/{conversation_id}/messages")
def stream_message(
    conversation_id: int,
    data: MessageRequest,
    request: Request,
):
    user_id = require_user(request)

    enforce_rate_limit(user_id)

    conversation = get_conversation(conversation_id, user_id)

    if not conversation:
        raise HTTPException(
            status_code=404,
            detail="Conversation not found.",
        )

    user_message = data.message.strip()
    attachments = validate_images(data.images)

    if not user_message and not attachments:
        raise HTTPException(
            status_code=400,
            detail="Please enter a message.",
        )

    if len(user_message) > 12000:
        raise HTTPException(
            status_code=400,
            detail="Message is too long (12,000 character limit).",
        )

    save_message(
        conversation_id, "user", user_message, attachments=attachments or None
    )
    rename_conversation_if_default(conversation_id, user_message)

    return build_streaming_response(
        conversation_id,
        user_id,
        data.language,
        data.research,
        data.image_aspect_ratio,
    )


@app.put("/api/conversations/{conversation_id}/messages/{message_id}")
def edit_message(
    conversation_id: int,
    message_id: int,
    data: MessageRequest,
    request: Request,
):
    """Edits a past user message: deletes it and everything after it,
    then re-asks with the new content and streams a fresh reply."""
    user_id = require_user(request)

    enforce_rate_limit(user_id)

    conversation = get_conversation(conversation_id, user_id)

    if not conversation:
        raise HTTPException(
            status_code=404,
            detail="Conversation not found.",
        )

    existing = get_message(conversation_id, message_id)

    if not existing or existing[1] != "user":
        raise HTTPException(
            status_code=404,
            detail="Message not found.",
        )

    new_content = data.message.strip()
    attachments = validate_images(data.images)

    if not new_content and not attachments:
        raise HTTPException(
            status_code=400,
            detail="Please enter a message.",
        )

    if len(new_content) > 12000:
        raise HTTPException(
            status_code=400,
            detail="Message is too long (12,000 character limit).",
        )

    delete_messages_from(conversation_id, message_id)
    save_message(
        conversation_id, "user", new_content, attachments=attachments or None
    )
    rename_conversation_if_default(conversation_id, new_content)

    return build_streaming_response(
        conversation_id,
        user_id,
        data.language,
        data.research,
        data.image_aspect_ratio,
    )


@app.post("/api/conversations/{conversation_id}/regenerate")
def regenerate_last(
    conversation_id: int,
    request: Request,
    data: RegenerateRequest = RegenerateRequest(),
):
    """Deletes the last assistant message (if any) and streams a fresh
    reply to the last remaining user message."""
    user_id = require_user(request)

    enforce_rate_limit(user_id)

    conversation = get_conversation(conversation_id, user_id)

    if not conversation:
        raise HTTPException(
            status_code=404,
            detail="Conversation not found.",
        )

    last_message = get_last_message(conversation_id)

    if not last_message:
        raise HTTPException(
            status_code=400,
            detail="Nothing to regenerate yet.",
        )

    last_id, last_role, _last_content, _last_attachments = last_message

    if last_role == "assistant":
        delete_messages_from(conversation_id, last_id)

    return build_streaming_response(
        conversation_id,
        user_id,
        data.language,
        data.research,
        data.image_aspect_ratio,
    )


# --- Routes: attachments (PDF / DOCX / PPTX text extraction) ----------------

@app.post("/api/extract-text")
async def extract_text(
    request: Request,
    file: UploadFile = File(...),
):
    """Extracts plain text from an uploaded PDF, DOCX, or PPTX file so
    the client can inline it into a message, the same way it already
    does for plain-text files. See extract_pdf_text() /
    extract_docx_text() / extract_pptx_text() for limits — scanned/
    image-only PDFs have no text layer to extract (OCR is a follow-up,
    not handled here), and legacy binary .ppt isn't supported (only
    OOXML .pptx)."""
    user_id = require_user(request)

    enforce_rate_limit(user_id)

    filename = file.filename or "document"
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if suffix not in ("pdf", "docx", "pptx"):
        raise HTTPException(
            status_code=400,
            detail="Only PDF, DOCX, and PPTX files are supported.",
        )

    data = await file.read()

    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=400,
            detail="File is too large (15 MB limit).",
        )

    try:
        if suffix == "pdf":
            text, truncated = extract_pdf_text(data)
        elif suffix == "docx":
            text, truncated = extract_docx_text(data)
        else:
            text, truncated = extract_pptx_text(data)

    except Exception as error:
        raise HTTPException(
            status_code=400,
            detail=f"Could not read that file: {error}",
        ) from error

    if not text.strip():
        raise HTTPException(
            status_code=422,
            detail=(
                "No extractable text was found in that file. If it's a "
                "scanned/image-only PDF, note that OCR isn't supported "
                "yet."
            ),
        )

    return {
        "filename": filename,
        "text": text,
        "truncated": truncated,
    }


# --- Routes: generation (documents, images, canvas) --------------------------

@app.get("/api/generated/{file_id}/download")
def download_generated_file(file_id: int, request: Request):
    user_id = require_user(request)

    row = get_generated_file(file_id, user_id)

    if not row:
        raise HTTPException(
            status_code=404,
            detail="File not found.",
        )

    title, _kind, mime_type, data = row

    extension = mimetypes.guess_extension(mime_type) or ".bin"
    safe_title = re.sub(r"[^A-Za-z0-9 _-]", "", title).strip() or "download"

    return Response(
        content=bytes(data),
        media_type=mime_type,
        headers={
            "Content-Disposition": (
                f'attachment; filename="{safe_title}{extension}"'
            ),
        },
    )


@app.post("/api/regenerate-image")
def regenerate_image(data: RegenerateImageRequest, request: Request):
    """Backs each inline image result's "Regenerate" button -- a
    lightweight direct re-run of the same prompt (optionally a new
    aspect ratio), independent of the chat/tool-calling loop, so
    regenerating doesn't add a new turn to the conversation transcript.
    Returns a new generated_files row; the client swaps the displayed
    image in place."""
    user_id = require_user(request)

    enforce_rate_limit(user_id)

    if not ENABLE_GENERATION:
        raise HTTPException(
            status_code=403,
            detail="Image generation is disabled.",
        )

    prompt = (data.prompt or "").strip()

    if not prompt:
        raise HTTPException(
            status_code=400,
            detail="No image prompt was provided.",
        )

    aspect_ratio = data.aspect_ratio if data.aspect_ratio in IMAGE_ASPECT_RATIOS else "square"

    try:
        image_bytes, mime_type, cost_usd = generate_image_bytes(prompt, aspect_ratio)
    except Exception as error:
        raise HTTPException(
            status_code=502,
            detail=f"Image generation failed: {error}",
        ) from error

    file_id = save_generated_file(
        user_id, data.conversation_id, prompt[:150], "image", mime_type, image_bytes
    )

    try:
        log_direct_cost(user_id, data.conversation_id, IMAGE_MODEL, cost_usd)
    except Exception as error:  # noqa: BLE001 - best-effort
        report_error("image-regeneration cost logging failed", error)

    return {
        "id": file_id,
        "prompt": prompt,
        "aspect_ratio": aspect_ratio,
        "download_url": f"/api/generated/{file_id}/download",
        "size_bytes": len(image_bytes),
    }


@app.get("/api/conversations/{conversation_id}/canvas")
def get_canvas(conversation_id: int, request: Request):
    user_id = require_user(request)

    conversation = get_conversation(conversation_id, user_id)

    if not conversation:
        raise HTTPException(
            status_code=404,
            detail="Conversation not found.",
        )

    artifact = get_canvas_artifact(conversation_id)

    if not artifact:
        return {"canvas": None}

    (
        artifact_id, title, kind, language, content,
        generated_file_id, updated_at,
    ) = artifact

    return {
        "canvas": {
            "id": artifact_id,
            "title": title,
            "kind": kind,
            "language": language,
            "content": content,
            "download_url": (
                f"/api/generated/{generated_file_id}/download"
                if generated_file_id else None
            ),
            "updated_at": serialize_datetime(updated_at),
        }
    }


# --- Routes: voice (server-side STT/TTS fallbacks) ---------------------------

@app.post("/api/transcribe")
async def transcribe(
    request: Request,
    file: UploadFile = File(...),
    conversation_id: int | None = Form(None),
):
    """Server-side speech-to-text fallback for when the browser has no
    (or an unreliable) SpeechRecognition implementation -- see
    ENABLE_SERVER_STT and the composer mic button in chat.js, which
    only calls this when the browser API is unavailable."""
    user_id = require_user(request)

    enforce_rate_limit(user_id)

    if not ENABLE_SERVER_STT:
        raise HTTPException(
            status_code=403,
            detail="Server-side transcription is disabled.",
        )

    data = await file.read()

    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=400,
            detail="Audio is too large (15 MB limit).",
        )

    audio_format = (file.content_type or "").split("/")[-1].split(";")[0]
    audio_format = audio_format or "webm"

    try:
        text, cost_usd = transcribe_audio_bytes(data, audio_format)
    except Exception as error:
        raise HTTPException(
            status_code=502,
            detail=f"Transcription failed: {error}",
        ) from error

    if cost_usd:
        try:
            log_direct_cost(user_id, conversation_id, STT_MODEL, cost_usd)
        except Exception as error:  # noqa: BLE001 - best-effort
            report_error("STT cost logging failed", error)

    return {"text": text}


@app.post("/api/speak")
def speak(data: SpeakRequest, request: Request):
    """Server-side text-to-speech fallback for when the browser has no
    (or a low-quality/wrong-language) speechSynthesis voice -- see
    ENABLE_SERVER_TTS and each assistant message's "Read aloud" button
    in chat.js, which only calls this when the browser can't cover the
    current reply language."""
    user_id = require_user(request)

    enforce_rate_limit(user_id)

    if not ENABLE_SERVER_TTS:
        raise HTTPException(
            status_code=403,
            detail="Server-side voice is disabled.",
        )

    text = (data.text or "").strip()

    if not text:
        raise HTTPException(
            status_code=400,
            detail="No text was provided.",
        )

    truncated = len(text) > MAX_TTS_CHARS
    text = text[:MAX_TTS_CHARS]

    try:
        audio_bytes, cost_usd = synthesize_speech_bytes(text)
    except Exception as error:
        raise HTTPException(
            status_code=502,
            detail=f"Voice synthesis failed: {error}",
        ) from error

    try:
        log_direct_cost(user_id, None, TTS_MODEL, cost_usd)
    except Exception as error:  # noqa: BLE001 - best-effort
        report_error("TTS cost logging failed", error)

    return Response(
        content=audio_bytes,
        media_type="audio/mpeg",
        headers={"X-Truncated": "true" if truncated else "false"},
    )


# --- Bulk audio transcription --------------------------------------------
# "Transcribe a large dataset of audios" is a batch job, not a single
# request/response -- upload returns a job id immediately, and
# transcription happens afterward. This first version uses FastAPI
# BackgroundTasks (same process, runs after the response is sent), which
# is fine for modest volume but has real limits worth being explicit
# about:
#   - No cross-instance durability: if the backend restarts or
#     redeploys mid-job (or runs behind more than one instance), a
#     queued/processing item can be stuck rather than picked up
#     elsewhere -- there's no separate worker process or persistent
#     queue watching for orphaned jobs.
#   - No backpressure: a user could enqueue many large jobs back to
#     back with only the per-file/per-job caps below limiting them,
#     no global concurrency cap across jobs.
# For genuinely large batches (hundreds of files / many hours of
# audio), swap this for a real task queue -- Celery or RQ backed by
# Redis, or a dedicated Sevalla background worker process -- consuming
# from the same transcription_items table (or a proper queue) with
# bounded concurrency per worker. See README "Suggested next steps".

MAX_FILES_PER_TRANSCRIPTION_JOB = 50
# OpenRouter's transcription endpoint documents a 25 MB multipart
# upload limit (~26 min of 128kbps MP3) -- enforced here too so an
# oversized file fails fast with a clear error instead of a confusing
# provider error later.
MAX_AUDIO_FILE_BYTES = 25 * 1024 * 1024
# Guards the overall upload (particularly a .zip) against abuse --
# well above what MAX_FILES_PER_TRANSCRIPTION_JOB legitimate files
# would need.
MAX_TRANSCRIPTION_UPLOAD_BYTES = 200 * 1024 * 1024
MAX_TRANSCRIPTION_CONCURRENCY = 3

SUPPORTED_AUDIO_EXTENSIONS = {
    "mp3", "mp4", "mpeg", "mpga", "m4a", "wav", "webm", "ogg", "flac",
}


def create_transcription_job(user_id: int, total_files: int) -> int:
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO transcription_jobs (user_id, total_files)
                VALUES (%s, %s)
                RETURNING id;
                """,
                (user_id, total_files),
            )

            return cursor.fetchone()[0]


def create_transcription_item(
    job_id: int,
    filename: str,
    data: bytes | None,
    status: str = "queued",
    error: str | None = None,
) -> int:
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                INSERT INTO transcription_items (
                    job_id, filename, data, status, error
                )
                VALUES (%s, %s, %s, %s, %s)
                RETURNING id;
                """,
                (job_id, filename[:255], data, status, error),
            )

            return cursor.fetchone()[0]


def get_transcription_job(job_id: int, user_id: int):
    """Returns (id, status, total_files, completed_files, created_at)
    or None."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT id, status, total_files, completed_files, created_at
                FROM transcription_jobs
                WHERE id = %s AND user_id = %s;
                """,
                (job_id, user_id),
            )

            return cursor.fetchone()


def get_transcription_job_owner(job_id: int) -> int | None:
    """Returns just the owning user_id for a job, with no ownership
    check -- used internally by the background worker (which already
    trusts job_id, having created the job itself), not by any route
    that takes job_id from a request."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT user_id FROM transcription_jobs WHERE id = %s;",
                (job_id,),
            )
            row = cursor.fetchone()
            return row[0] if row else None


def get_transcription_jobs(user_id: int):
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT id, status, total_files, completed_files, created_at
                FROM transcription_jobs
                WHERE user_id = %s
                ORDER BY created_at DESC
                LIMIT 100;
                """,
                (user_id,),
            )

            return cursor.fetchall()


def get_transcription_items(job_id: int):
    """Returns rows of (id, filename, status, transcript_text, error,
    created_at) -- never `data` (the raw audio bytes)."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT id, filename, status, transcript_text, error, created_at
                FROM transcription_items
                WHERE job_id = %s
                ORDER BY id;
                """,
                (job_id,),
            )

            return cursor.fetchall()


def get_transcription_items_to_process(job_id: int):
    """Returns rows of (id, filename, data) for items still queued --
    used by the background worker, not exposed to the API."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT id, filename, data
                FROM transcription_items
                WHERE job_id = %s AND status = 'queued';
                """,
                (job_id,),
            )

            return cursor.fetchall()


def update_transcription_job_status(job_id: int, status: str):
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE transcription_jobs
                SET status = %s
                WHERE id = %s;
                """,
                (status, job_id),
            )


def increment_transcription_job_completed(job_id: int):
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE transcription_jobs
                SET completed_files = completed_files + 1
                WHERE id = %s;
                """,
                (job_id,),
            )


def update_transcription_item_status(item_id: int, status: str):
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE transcription_items
                SET status = %s
                WHERE id = %s;
                """,
                (status, item_id),
            )


def save_transcription_item_result(
    item_id: int,
    status: str,
    transcript_text: str | None = None,
    error: str | None = None,
):
    """Sets the item's final status (done/failed) and clears `data`
    (the raw audio bytes are no longer needed once processed, so
    clearing them bounds storage growth)."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE transcription_items
                SET status = %s, transcript_text = %s, error = %s, data = NULL
                WHERE id = %s;
                """,
                (status, transcript_text, error, item_id),
            )


def get_transcription_item_with_job(item_id: int, job_id: int, user_id: int):
    """Returns (filename, status, transcript_text, error) for an item,
    scoped to a job owned by user_id, or None."""
    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT ti.filename, ti.status, ti.transcript_text, ti.error
                FROM transcription_items ti
                JOIN transcription_jobs tj ON tj.id = ti.job_id
                WHERE ti.id = %s AND ti.job_id = %s AND tj.user_id = %s;
                """,
                (item_id, job_id, user_id),
            )

            return cursor.fetchone()


def extract_audio_entries_from_zip(zip_bytes: bytes) -> list[tuple[str, bytes]]:
    """Returns (filename, data) for every non-directory entry in a
    .zip upload -- including files with an unrecognized extension, so
    they can surface as a clear "unsupported format" failed item
    rather than being silently skipped during extraction."""
    entries = []

    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as archive:
        for info in archive.infolist():
            if info.is_dir():
                continue

            if len(entries) >= MAX_FILES_PER_TRANSCRIPTION_JOB:
                break

            entries.append((info.filename, archive.read(info)))

    return entries


def process_transcription_job(job_id: int):
    """Runs after the upload response has already been sent (FastAPI
    BackgroundTasks) -- transcribes every queued item in the job, with
    limited concurrency (MAX_TRANSCRIPTION_CONCURRENCY), and marks the
    job 'done' when finished. See the module comment above this
    section for why this doesn't scale to hundreds of files."""
    update_transcription_job_status(job_id, "processing")

    owner_user_id = get_transcription_job_owner(job_id)
    items = get_transcription_items_to_process(job_id)

    def process_one(item):
        item_id, filename, data = item

        update_transcription_item_status(item_id, "processing")

        try:
            extension = (
                filename.rsplit(".", 1)[-1].lower()
                if "." in filename else ""
            )

            if extension not in SUPPORTED_AUDIO_EXTENSIONS:
                raise ValueError(
                    f"Unsupported format '{extension or 'unknown'}'."
                )

            if not data or len(data) > MAX_AUDIO_FILE_BYTES:
                raise ValueError(
                    "File exceeds the 25 MB provider limit."
                )

            text, cost_usd = transcribe_audio_bytes(data, extension)

            if not text.strip():
                raise ValueError("No speech was detected in this file.")

            save_transcription_item_result(item_id, "done", transcript_text=text)

            if cost_usd and owner_user_id is not None:
                try:
                    log_direct_cost(owner_user_id, None, STT_MODEL, cost_usd)
                except Exception as usage_error:  # noqa: BLE001 - best-effort
                    print(
                        "EASTA: bulk transcription cost logging failed: "
                        f"{usage_error!r}"
                    )

        except Exception as error:  # noqa: BLE001 - a per-item failure shouldn't stop the job
            save_transcription_item_result(
                item_id, "failed", error=str(error)
            )

        increment_transcription_job_completed(job_id)

    with concurrent.futures.ThreadPoolExecutor(
        max_workers=MAX_TRANSCRIPTION_CONCURRENCY
    ) as executor:
        list(executor.map(process_one, items))

    update_transcription_job_status(job_id, "done")


# --- Routes: bulk audio transcription ----------------------------------

@app.post("/api/transcription-jobs")
async def create_transcription_job_route(
    request: Request,
    background_tasks: BackgroundTasks,
    files: list[UploadFile] = File(...),
):
    """Accepts one or more audio files, or a .zip of them, in a single
    request; creates a job + one item per file, and returns
    immediately -- transcription runs in the background (see
    process_transcription_job()). Every item (including ones that will
    fail validation) is created as 'queued' and validated uniformly
    during processing, so upload itself stays fast regardless of batch
    size."""
    user_id = require_user(request)

    enforce_rate_limit(user_id)

    if not ENABLE_SERVER_STT:
        raise HTTPException(
            status_code=403,
            detail="Server-side transcription is disabled.",
        )

    entries: list[tuple[str, bytes]] = []
    total_bytes = 0

    for upload in files:
        data = await upload.read()
        total_bytes += len(data)

        if total_bytes > MAX_TRANSCRIPTION_UPLOAD_BYTES:
            raise HTTPException(
                status_code=400,
                detail="Upload is too large (200 MB limit).",
            )

        filename = upload.filename or "audio"

        if filename.lower().endswith(".zip"):
            entries.extend(extract_audio_entries_from_zip(data))
        else:
            entries.append((filename, data))

    if not entries:
        raise HTTPException(
            status_code=400,
            detail="No audio files were found in the upload.",
        )

    truncated = len(entries) > MAX_FILES_PER_TRANSCRIPTION_JOB
    entries = entries[:MAX_FILES_PER_TRANSCRIPTION_JOB]

    job_id = create_transcription_job(user_id, len(entries))

    for filename, data in entries:
        create_transcription_item(job_id, filename, data)

    background_tasks.add_task(process_transcription_job, job_id)

    return {
        "job_id": job_id,
        "total_files": len(entries),
        "truncated": truncated,
    }


@app.get("/api/transcription-jobs")
def list_transcription_jobs(request: Request):
    user_id = require_user(request)

    rows = get_transcription_jobs(user_id)

    return {
        "jobs": [
            {
                "id": row[0],
                "status": row[1],
                "total_files": row[2],
                "completed_files": row[3],
                "created_at": serialize_datetime(row[4]),
            }
            for row in rows
        ]
    }


@app.get("/api/transcription-jobs/{job_id}")
def get_transcription_job_route(job_id: int, request: Request):
    user_id = require_user(request)

    job = get_transcription_job(job_id, user_id)

    if not job:
        raise HTTPException(
            status_code=404,
            detail="Transcription job not found.",
        )

    items = get_transcription_items(job_id)

    return {
        "job": {
            "id": job[0],
            "status": job[1],
            "total_files": job[2],
            "completed_files": job[3],
            "created_at": serialize_datetime(job[4]),
        },
        "items": [
            {
                "id": row[0],
                "filename": row[1],
                "status": row[2],
                "has_transcript": bool(row[3]),
                "error": row[4],
                "created_at": serialize_datetime(row[5]),
            }
            for row in items
        ],
    }


@app.get("/api/transcription-jobs/{job_id}/items/{item_id}")
def get_transcription_item_route(job_id: int, item_id: int, request: Request):
    user_id = require_user(request)

    item = get_transcription_item_with_job(item_id, job_id, user_id)

    if not item:
        raise HTTPException(
            status_code=404,
            detail="Transcription item not found.",
        )

    filename, status, transcript_text, error = item

    return {
        "filename": filename,
        "status": status,
        "transcript_text": transcript_text,
        "error": error,
    }


@app.get("/api/transcription-jobs/{job_id}/download")
def download_transcription_job(
    job_id: int,
    request: Request,
    format: str = "txt",
):
    user_id = require_user(request)

    job = get_transcription_job(job_id, user_id)

    if not job:
        raise HTTPException(
            status_code=404,
            detail="Transcription job not found.",
        )

    items = get_transcription_items(job_id)
    done_items = [row for row in items if row[2] == "done" and row[3]]

    if not done_items:
        raise HTTPException(
            status_code=400,
            detail="No completed transcripts to download yet.",
        )

    if format == "zip":
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            for _id, filename, _status, transcript_text, _error, _created_at in done_items:
                base_name = filename.rsplit(".", 1)[0] if "." in filename else filename
                docx_bytes = render_markdown_to_docx(base_name, transcript_text)
                archive.writestr(f"{base_name}.docx", docx_bytes)

        return Response(
            content=buffer.getvalue(),
            media_type="application/zip",
            headers={
                "Content-Disposition": (
                    f'attachment; filename="transcription-job-{job_id}.zip"'
                ),
            },
        )

    # Default: one concatenated .txt file, each transcript under a
    # filename header.
    parts = [
        f"=== {filename} ===\n{transcript_text}"
        for _id, filename, _status, transcript_text, _error, _created_at in done_items
    ]
    combined_text = "\n\n".join(parts)

    return Response(
        content=combined_text.encode("utf-8"),
        media_type="text/plain",
        headers={
            "Content-Disposition": (
                f'attachment; filename="transcription-job-{job_id}.txt"'
            ),
        },
    )


# --- Routes: documents (RAG knowledge base) ---------------------------------

def _validate_document_fields(data: DocumentRequest) -> tuple[str, str]:
    title = data.title.strip()[:200] or "Untitled"
    content = data.content.strip()

    if not content:
        raise HTTPException(
            status_code=400,
            detail="Document content is required.",
        )

    if len(content) > 200_000:
        raise HTTPException(
            status_code=400,
            detail="Document is too large (200,000 character limit).",
        )

    return title, content


def _serialize_document_rows(rows) -> list[dict]:
    return [
        {
            "id": row[0],
            "title": row[1],
            "created_at": serialize_datetime(row[2]),
            "length": row[3],
        }
        for row in rows
    ]


@app.get("/api/documents")
def list_documents(request: Request):
    """Personal knowledge base only -- see the
    /api/organizations/{org_id}/documents routes below for an org's
    shared one."""
    user_id = require_user(request)

    return {"documents": _serialize_document_rows(get_document_rows(user_id=user_id))}


@app.post("/api/documents")
def add_document(data: DocumentRequest, request: Request):
    user_id = require_user(request)

    title, content = _validate_document_fields(data)
    row = create_document_row(title, content, user_id=user_id)

    return {
        "document": {
            "id": row[0],
            "title": row[1],
            "created_at": serialize_datetime(row[2]),
        }
    }


@app.delete("/api/documents/{document_id}")
def delete_document(document_id: int, request: Request):
    user_id = require_user(request)

    if not delete_document_row(document_id, user_id=user_id):
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    return {"message": "Document deleted."}


# --- Routes: organizations (Phase 22) ---------------------------------------

@app.post("/api/organizations")
def create_organization_route(data: CreateOrganizationRequest, request: Request):
    user_id = require_user(request)

    name = data.name.strip()[:100]

    if not name:
        raise HTTPException(
            status_code=400,
            detail="Organization name is required.",
        )

    org = create_organization(name, user_id)

    return {
        "organization": {
            "id": org[0],
            "name": org[1],
            "invite_url": f"{FRONTEND_URL}/join/{org[2]}",
            "created_at": serialize_datetime(org[3]),
            "role": "owner",
        }
    }


@app.get("/api/organizations")
def list_organizations_route(request: Request):
    user_id = require_user(request)

    rows = get_user_organizations(user_id)

    return {
        "organizations": [
            {"id": org_id, "name": name, "role": role}
            for org_id, name, role in rows
        ]
    }


@app.get("/api/organizations/{org_id}")
def get_organization_route(org_id: int, request: Request):
    user_id = require_user(request)
    role = require_org_member(org_id, user_id)

    org = get_organization(org_id)

    if not org:
        raise HTTPException(status_code=404, detail="Organization not found.")

    return {
        "organization": {
            "id": org[0],
            "name": org[1],
            "invite_url": f"{FRONTEND_URL}/join/{org[2]}",
            "created_at": serialize_datetime(org[3]),
            "role": role,
        }
    }


@app.post("/api/organizations/{org_id}/invite/regenerate")
def regenerate_invite_route(org_id: int, request: Request):
    user_id = require_user(request)
    require_org_owner(org_id, user_id)

    new_token = regenerate_organization_invite_token(org_id)

    return {"invite_url": f"{FRONTEND_URL}/join/{new_token}"}


@app.post("/api/organizations/join/{invite_token}")
def join_organization_route(invite_token: str, request: Request):
    user_id = require_user(request)

    org = get_organization_by_invite_token(invite_token)

    if not org:
        raise HTTPException(
            status_code=404,
            detail="This invite link is invalid or no longer active.",
        )

    join_organization(org[0], user_id, role="member")

    return {"organization": {"id": org[0], "name": org[1]}}


@app.get("/api/organizations/{org_id}/members")
def list_organization_members_route(org_id: int, request: Request):
    user_id = require_user(request)
    require_org_member(org_id, user_id)

    rows = get_organization_members(org_id)

    return {
        "members": [
            {
                "user_id": member_user_id,
                "username": username,
                "email": email,
                "role": role,
                "joined_at": serialize_datetime(joined_at),
            }
            for member_user_id, username, email, role, joined_at in rows
        ]
    }


@app.delete("/api/organizations/{org_id}/members/{member_user_id}")
def remove_organization_member_route(
    org_id: int, member_user_id: int, request: Request
):
    user_id = require_user(request)
    require_org_owner(org_id, user_id)

    target_role = get_organization_membership(org_id, member_user_id)

    if not target_role:
        raise HTTPException(status_code=404, detail="Member not found.")

    if target_role == "owner" and count_organization_owners(org_id) <= 1:
        raise HTTPException(
            status_code=400,
            detail=(
                "Can't remove the only owner -- promote another member "
                "to owner first."
            ),
        )

    remove_organization_member(org_id, member_user_id)

    return {"message": "Member removed."}


@app.post("/api/organizations/{org_id}/leave")
def leave_organization_route(org_id: int, request: Request):
    user_id = require_user(request)
    role = require_org_member(org_id, user_id)

    if role == "owner" and count_organization_owners(org_id) <= 1:
        raise HTTPException(
            status_code=400,
            detail=(
                "You're the only owner -- promote another member to "
                "owner before leaving."
            ),
        )

    remove_organization_member(org_id, user_id)

    return {"message": "Left the organization."}


@app.get("/api/organizations/{org_id}/usage")
def organization_usage_route(org_id: int, request: Request):
    """Combined cost dashboard -- owner-only (not every member), same
    as the brief specifies. Deliberately separate from GET
    /api/usage/summary, which stays scoped to the caller's own data
    regardless of org membership."""
    user_id = require_user(request)
    require_org_owner(org_id, user_id)

    rows = get_organization_usage(org_id)

    members = [
        {
            "user_id": member_user_id,
            "username": username,
            "total_cost": float(total_cost),
            "call_count": call_count,
        }
        for member_user_id, username, total_cost, call_count in rows
    ]

    return {
        "total_cost": sum(member["total_cost"] for member in members),
        "total_calls": sum(member["call_count"] for member in members),
        "members": members,
    }


@app.get("/api/organizations/{org_id}/documents")
def list_organization_documents_route(org_id: int, request: Request):
    user_id = require_user(request)
    require_org_member(org_id, user_id)

    return {
        "documents": _serialize_document_rows(get_document_rows(org_id=org_id))
    }


@app.post("/api/organizations/{org_id}/documents")
def add_organization_document_route(
    org_id: int, data: DocumentRequest, request: Request
):
    user_id = require_user(request)
    require_org_member(org_id, user_id)

    title, content = _validate_document_fields(data)
    row = create_document_row(title, content, org_id=org_id)

    return {
        "document": {
            "id": row[0],
            "title": row[1],
            "created_at": serialize_datetime(row[2]),
        }
    }


@app.delete("/api/organizations/{org_id}/documents/{document_id}")
def delete_organization_document_route(
    org_id: int, document_id: int, request: Request
):
    user_id = require_user(request)
    require_org_member(org_id, user_id)

    if not delete_document_row(document_id, org_id=org_id):
        raise HTTPException(
            status_code=404,
            detail="Document not found.",
        )

    return {"message": "Document deleted."}


# --- Routes: admin (Phase 23) ------------------------------------------------
# Platform-wide, instance-owner-only visibility -- distinct from an
# organization owner's combined dashboard above (that's scoped to one
# org; this is every user/org/conversation on the whole instance). See
# require_admin() and the is_admin column comment in database/schema.sql
# for how someone gets this -- there's no self-service route, on purpose.

@app.get("/api/admin/stats")
def admin_stats_route(request: Request):
    user_id = require_user(request)
    require_admin(user_id)

    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute("SELECT COUNT(*) FROM users;")
            total_users = cursor.fetchone()[0]

            cursor.execute("SELECT COUNT(*) FROM organizations;")
            total_organizations = cursor.fetchone()[0]

            cursor.execute("SELECT COUNT(*) FROM conversations;")
            total_conversations = cursor.fetchone()[0]

            cursor.execute("SELECT COUNT(*) FROM messages;")
            total_messages = cursor.fetchone()[0]

            cursor.execute("SELECT COUNT(*) FROM documents;")
            total_documents = cursor.fetchone()[0]

            cursor.execute(
                "SELECT COALESCE(SUM(cost_usd), 0), COUNT(*) FROM usage_logs;"
            )
            total_cost_usd, total_calls = cursor.fetchone()

            cursor.execute(
                """
                SELECT
                    date_trunc('day', created_at)::date AS day,
                    COUNT(*)
                FROM users
                WHERE created_at > NOW() - INTERVAL '30 days'
                GROUP BY day
                ORDER BY day;
                """
            )
            signups_by_day = cursor.fetchall()

    return {
        "total_users": total_users,
        "total_organizations": total_organizations,
        "total_conversations": total_conversations,
        "total_messages": total_messages,
        "total_documents": total_documents,
        "total_cost_usd": float(total_cost_usd),
        "total_calls": total_calls,
        "signups_by_day": [
            {"day": row[0].isoformat(), "count": row[1]}
            for row in signups_by_day
        ],
    }


@app.get("/api/admin/users")
def admin_users_route(request: Request):
    user_id = require_user(request)
    require_admin(user_id)

    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            # LEFT JOINs so a brand-new user with zero conversations/
            # spend still shows up with 0 rather than being silently
            # absent -- same reasoning as get_organization_usage()'s
            # per-member breakdown above.
            cursor.execute(
                """
                SELECT
                    u.id,
                    u.username,
                    u.email,
                    u.plan,
                    u.is_admin,
                    u.created_at,
                    COUNT(DISTINCT c.id) AS conversation_count,
                    COALESCE(SUM(ul.cost_usd), 0) AS total_cost
                FROM users u
                LEFT JOIN conversations c ON c.user_id = u.id
                LEFT JOIN usage_logs ul ON ul.user_id = u.id
                GROUP BY u.id
                ORDER BY u.created_at DESC;
                """
            )
            rows = cursor.fetchall()

    return {
        "users": [
            {
                "id": row[0],
                "username": row[1],
                "email": row[2],
                "plan": row[3],
                "is_admin": row[4],
                "created_at": serialize_datetime(row[5]),
                "conversation_count": row[6],
                "total_cost": float(row[7]),
            }
            for row in rows
        ]
    }


# --- Routes: usage / cost dashboard -----------------------------------------

@app.get("/api/usage/summary")
def usage_summary(request: Request):
    user_id = require_user(request)

    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT
                    COALESCE(SUM(cost_usd), 0),
                    COALESCE(SUM(prompt_tokens), 0),
                    COALESCE(SUM(completion_tokens), 0),
                    COUNT(*)
                FROM usage_logs
                WHERE user_id = %s;
                """,
                (user_id,),
            )
            totals = cursor.fetchone()

            cursor.execute(
                """
                SELECT
                    model,
                    SUM(cost_usd),
                    SUM(prompt_tokens + completion_tokens),
                    COUNT(*)
                FROM usage_logs
                WHERE user_id = %s
                GROUP BY model
                ORDER BY SUM(cost_usd) DESC;
                """,
                (user_id,),
            )
            by_model = cursor.fetchall()

            cursor.execute(
                """
                SELECT
                    date_trunc('day', created_at)::date AS day,
                    SUM(cost_usd)
                FROM usage_logs
                WHERE
                    user_id = %s
                    AND created_at > NOW() - INTERVAL '30 days'
                GROUP BY day
                ORDER BY day;
                """,
                (user_id,),
            )
            by_day = cursor.fetchall()

    return {
        "total_cost_usd": float(totals[0]),
        "total_prompt_tokens": totals[1],
        "total_completion_tokens": totals[2],
        "total_calls": totals[3],
        "by_model": [
            {
                "model": row[0],
                "cost_usd": float(row[1] or 0),
                "tokens": row[2] or 0,
                "calls": row[3],
            }
            for row in by_model
        ],
        "by_day": [
            {
                "day": row[0].isoformat(),
                "cost_usd": float(row[1] or 0),
            }
            for row in by_day
        ],
    }


@app.get("/api/usage/logs")
def usage_logs_endpoint(request: Request, limit: int = 50):
    user_id = require_user(request)
    limit = max(1, min(limit, 200))

    with psycopg.connect(DATABASE_URL) as connection:
        with connection.cursor() as cursor:
            cursor.execute(
                """
                SELECT
                    model,
                    prompt_tokens,
                    completion_tokens,
                    cost_usd,
                    created_at
                FROM usage_logs
                WHERE user_id = %s
                ORDER BY created_at DESC
                LIMIT %s;
                """,
                (user_id, limit),
            )
            rows = cursor.fetchall()

    return {
        "logs": [
            {
                "model": row[0],
                "prompt_tokens": row[1],
                "completion_tokens": row[2],
                "cost_usd": float(row[3]),
                "created_at": serialize_datetime(row[4]),
            }
            for row in rows
        ]
    }


# --- Public API (v1) ----------------------------------------------------
# A small, deliberately-scoped surface for third-party/programmatic
# access -- see API.md for docs + a curl example. Authenticated with an
# API key (Account page -> API Keys), via the same require_user() /
# enforce_rate_limit() every internal endpoint already uses, so this
# gets Bearer-token auth and per-user rate limiting for free rather
# than duplicating either.

def run_public_chat_turn(
    conversation_id: int,
    user_id: int,
    language: str | None = None,
) -> str:
    """Runs one assistant turn and returns the complete reply as a
    plain string -- the non-streaming counterpart to
    build_streaming_response(), used by POST /v1/chat. Deliberately
    simpler than the internal endpoint (no research mode, canvas, or
    image-style composer hints for v1 -- just RAG grounding, a
    reply-language override, and the same tool-calling loop via
    stream_with_tools()) rather than exposing every internal feature
    through the public API. This does duplicate some of
    build_streaming_response()'s turn logic rather than sharing it
    directly -- deliberate, to keep the streaming endpoint that the
    whole live chat UI depends on untouched by v1's simpler needs."""
    conversation = get_conversation(conversation_id, user_id)

    if not conversation:
        raise HTTPException(
            status_code=404,
            detail="Conversation not found.",
        )

    last_message = get_last_message(conversation_id)
    is_latest_from_user = bool(last_message) and last_message[1] == "user"
    latest_user_text = last_message[2] if is_latest_from_user else ""
    latest_user_attachments = (
        last_message[3] if is_latest_from_user else None
    )
    has_image_attachments = any(
        attachment.get("type") == "image"
        for attachment in (latest_user_attachments or [])
    )

    rag_message = build_rag_system_message(latest_user_text, user_id)
    language_message = build_language_override_message(language)
    canvas_message = build_canvas_context_message(conversation_id)

    llm_context = build_llm_context(
        conversation_id,
        conversation,
        rag_message,
        user_id=user_id,
        language_message=language_message,
        canvas_message=canvas_message,
    )

    models_to_try = (
        model_chain_for_images() if has_image_attachments
        else model_chain_for(latest_user_text)
    )

    complete_response = ""

    for attempt, model in enumerate(models_to_try):
        sent_any_output = False

        try:
            for kind, payload in stream_with_tools(
                llm_context, model, user_id=user_id, conversation_id=conversation_id,
            ):
                if kind == "token":
                    complete_response += payload
                    sent_any_output = True

                elif kind == "canvas":
                    if payload.get("type") == "file_card":
                        extension = {
                            "pdf": ".pdf", "docx": ".docx", "pptx": ".pptx",
                        }.get(payload.get("format"), "")
                        complete_response += (
                            f"\n\n📄 [Download "
                            f"{payload['title']}{extension}]"
                            f"({payload['download_url']})"
                        )
                    elif payload.get("type") == "image_result":
                        complete_response += (
                            f"\n\n![{payload['prompt']}]"
                            f"({payload['download_url']})"
                        )

                elif kind == "usage":
                    try:
                        log_usage(
                            user_id,
                            conversation_id,
                            model,
                            payload.prompt_tokens or 0,
                            payload.completion_tokens or 0,
                        )
                    except Exception as usage_error:  # noqa: BLE001
                        report_error("usage logging failed", usage_error)

            break

        except Exception as error:
            report_error(
                f"model '{model}' failed (v1 chat, attempt "
                f"{attempt + 1}/{len(models_to_try)})",
                error,
            )

            if sent_any_output:
                break

            if attempt == len(models_to_try) - 1:
                raise HTTPException(
                    status_code=502,
                    detail="EASTA could not complete the response right now.",
                ) from error

            continue

    if complete_response:
        save_message(conversation_id, "assistant", complete_response)

        try:
            fresh_conversation = get_conversation(conversation_id, user_id)
            if fresh_conversation:
                maybe_summarize_conversation(conversation_id, fresh_conversation)
        except Exception as error:  # noqa: BLE001
            report_error("post-turn summarization skipped (v1 chat)", error)

        try:
            maybe_extract_memory(conversation_id, user_id, latest_user_text)
        except Exception as error:  # noqa: BLE001
            report_error("post-turn memory extraction skipped (v1 chat)", error)

    return complete_response


@app.post("/v1/chat")
def public_chat(data: PublicChatRequest, request: Request):
    user_id = require_user(request)

    enforce_rate_limit(user_id)

    message = data.message.strip()

    if not message:
        raise HTTPException(
            status_code=400,
            detail="message is required.",
        )

    if len(message) > 12000:
        raise HTTPException(
            status_code=400,
            detail="message is too long (12,000 character limit).",
        )

    if data.conversation_id is not None:
        conversation = get_conversation(data.conversation_id, user_id)

        if not conversation:
            raise HTTPException(
                status_code=404,
                detail="conversation_id not found.",
            )

        conversation_id = data.conversation_id
    else:
        conversation = create_conversation(user_id)
        conversation_id = conversation[0]

    save_message(conversation_id, "user", message)
    rename_conversation_if_default(conversation_id, message)

    reply_text = run_public_chat_turn(conversation_id, user_id, data.language)

    return {
        "conversation_id": conversation_id,
        "message": {
            "role": "assistant",
            "content": reply_text,
        },
    }


@app.get("/v1/me")
def public_me(request: Request):
    user_id = require_user(request)

    user = get_user_by_id(user_id)

    if not user:
        raise HTTPException(
            status_code=404,
            detail="Account not found.",
        )

    _id, username, _email, _password_hash, _created_at, plan, _is_admin = user

    return {"username": username, "plan": plan}
